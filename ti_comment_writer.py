"""EHT비교 시트에 적은 COMMENT 를 개별 TI-LST 파일에 옮겨 적는다.

사용법
  1. ti_lst_reader.py 결과 파일의 'EHT비교' 시트 COMMENT 열(AC)에 코멘트를 적고,
     파일명에 COMMENT 가 들어가게 저장한다 (예: EHT비교_COMMENT.xlsx, 그 시트만 따로 저장해도 됨).
  2. 그 파일을 TI-LST 파일들이 있는 폴더(ti_lst_reader.py 와 같은 폴더)에 넣고 이 프로그램을 실행한다.
     (COMMENT 파일을 이 프로그램 위로 끌어다 놓아도 됨)
  3. 'OUTPUT TI COMMENT' 폴더에 COMMENT 가 들어간 TI-LST 가 '원래이름_COMMENT.xlsx' 로 저장된다.
     원본 TI-LST 는 바꾸지 않는다. 어디에 적었는지는 같은 폴더의 COMMENT_반영결과_*.xlsx 에 남는다.

어디에 적나
  - EHT비교에서 TI 와 비교한 행: 비교한 TI 품목(같은 ISO 시트·DN 의 파이프·밸브)의 품목 행 O열
  - 붙일 TI 품목이 없는 행(TI에 ISO 없음 등): 그 ISO 시트가 있는 TI-LST, 없으면 같은 구역에서 ISO 번호가
    가장 가까운 TI-LST 의 데이터가 끝난 다음 행부터 O열에 COMMENT 를 차례로 적음
  - COMMENT 에 두께(예: 'thickness shall be 50mm')가 하나만 적혀 있으면 이미 그 두께인 품목에는 적지 않음

결과 파일 모양
  - O열 머리글 'Comment' (N열 Note 머리글과 같은 서식), COMMENT 행은 노란색 음영
  - COMMENT 가 있는 행만 보이게 필터를 건 채로 저장 (필터를 풀면 전체가 보임), '기본' 보기
  원본 파일의 다른 내용(서식, 그림, 머리글/바닥글 등)은 그대로 둔다.
"""
from pathlib import Path
import os
import posixpath
import re
import sys
import time
import traceback
import xml.etree.ElementTree as ET
import zipfile
from collections import defaultdict
from datetime import datetime
from multiprocessing import freeze_support
from xml.parsers.expat import ParserCreate

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries

import ti_lst_reader as ti  # 같은 폴더의 ti_lst_reader.py (TI-LST 읽기·비교 규칙을 그대로 씀)


# =========================
# 사용자 설정
# =========================
COMMENT_FILE_KEYWORD = "COMMENT"     # 파일명에 이 글자가 있는 엑셀(EHT비교 시트)에서 COMMENT 를 읽음
OUTPUT_FOLDER = "OUTPUT TI COMMENT"  # 결과 폴더 (이 프로그램이 있는 폴더 안)
OUTPUT_SUFFIX = "_COMMENT"           # 결과 파일명: 원래이름_COMMENT.xlsx
COMMENT_COLUMN = "O"                 # TI-LST 에 COMMENT 를 적을 열 (N열 Note 바로 오른쪽)
COMMENT_HEADER = "Comment"           # COMMENT 열 머리글 (N열 머리글과 같은 서식). "" 이면 머리글 안 적음
COMMENT_COLUMN_WIDTH = 40            # COMMENT 열 너비. None 이면 원본 너비 그대로
HIGHLIGHT_COLOR = "FFFFFF00"         # COMMENT 행(A~O열) 음영 색 (ARGB, 노란색). None 이면 음영 없음
FILTER_COMMENT_ROWS = True           # True: COMMENT 가 있는 행만 보이게 필터를 건 채로 저장
NORMAL_VIEW = True                   # True: '기본' 보기로 저장 (페이지 나누기 미리 보기를 풀어 보기 쉽게)
SKIP_SAME_THICKNESS = True           # True: COMMENT 의 두께(예: 50mm)와 이미 같은 두께인 품목에는 안 적음
APPEND_FORMAT = "{comment}"          # 데이터 끝에 추가하는 COMMENT 형식 (예: "{iso} / {line} / DN {dn}: {comment}")


REPORT_KEY_COLUMNS = ("ISO NO", "DN", "TI LINE NO", "TI 파일")  # EHT비교 시트인지 확인하는 머리글
ISO_SERIAL_RE = re.compile(r"^(.*-ISO)-(\d+)(?:-\d{2})?$")     # ...-TKM-ISO-00007-01 -> (...-TKM-ISO, 7)

LOG_COLUMNS = [
    "결과", "COMMENT", "TI 파일", "TI 행", "품명", "Size (mm)", "Nominal Diameter", "두께 (mm)",
    "LINE NO", "ISO NO", "DN", "판정", "EHT 라인", "COMMENT 파일", "COMMENT 행", "저장 파일",
]


# =========================
# COMMENT 읽기
# =========================
def comment_text(value):
    if value is None:
        return ""
    return str(value).strip()


def read_comment_file(path):
    """EHT비교 시트(머리글로 찾음)에서 COMMENT 가 적힌 행을 읽는다. EHT비교 시트가 없으면 None.

    COMMENT 열: 머리글에 COMMENT 가 들어간 열, 없으면 'EHT 행' 다음 열(AC)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            ws.reset_dimensions()
            rows = ws.iter_rows(values_only=True)
            header = [ti.clean_text(v) for v in next(rows, ())]
            if not all(col in header for col in REPORT_KEY_COLUMNS):
                continue
            comment_col = next((i for i, h in enumerate(header) if "comment" in h.lower()), None)
            if comment_col is None:
                comment_col = header.index("EHT 행") + 1 if "EHT 행" in header else len(header)
            columns = {h: i for i, h in enumerate(header) if h and i != comment_col}

            records = []
            for row_no, row in enumerate(rows, start=2):
                text = comment_text(row[comment_col] if comment_col < len(row) else None)
                if not text:
                    continue
                rec = {h: (row[i] if i < len(row) else None) for h, i in columns.items()}
                rec.update({"COMMENT": text, "COMMENT 파일": path.name, "COMMENT 행": row_no})
                records.append(rec)
            return records, ws.title, get_column_letter(comment_col + 1)
    finally:
        wb.close()
    return None


def split_names(value, known=()):
    """'a, b' -> ['a', 'b'] (통째로 아는 이름이면 그대로)"""
    text = ti.clean_text(value)
    if not text:
        return []
    if text in known:
        return [text]
    return [part.strip() for part in text.split(", ") if part.strip()]


# =========================
# COMMENT 를 적을 TI 품목 찾기 (ti_lst_reader 의 EHT 비교 규칙 그대로)
# =========================
def comment_iso(rec):
    iso = ti.normalize_iso(rec.get("ISO NO"))
    if iso and not ti.has_sheet_no(iso):  # EHT ISO 에 시트번호가 없으면 라인번호 끝 숫자로 보충 (비교 때와 같음)
        sheet_no = ti.get_sheet_no(rec.get("EHT 라인"))
        if sheet_no:
            iso = ti.normalize_iso(f"{iso}-{sheet_no}")
    return iso


def comment_targets(rec, groups, known_files):
    """EHT비교 한 행에서 비교했던 TI 품목들.

    EHT비교의 TI LINE NO / TI 보온구분 / TI 파일 은 비교에 쓴 품목에서 나온 값이므로 그 값으로 다시 고르고,
    비교 때처럼 제외(Excluded) 라인은 빼고(그것만 있으면 그대로), 전기히팅 표시 품목이 있으면 그 품목만 쓴다."""
    entries = groups.get((comment_iso(rec), ti.dn_number(rec.get("DN"))), [])
    files = set(split_names(rec.get("TI 파일"), known_files))
    lines = set(split_names(rec.get("TI LINE NO")))
    types = set(split_names(rec.get("TI 보온구분")))
    picked = [e for e in entries
              if e["file"] in files and e["line"] in lines and (not types or e["type"] in types)]
    active = [e for e in picked if not e["excluded"]]
    if active:
        picked = [e for e in active if e["heating"]] or active
    return picked


def ti_groups(items):
    """ti_lst_reader 의 비교용 묶음 {(ISO 시트, DN): [TI 품목]}. 품목마다 원본 행 번호('row')를 붙여 돌려준다.

    예전 ti_lst_reader.py 는 'row' 를 넣지 않으므로, 같은 순서·같은 조건으로 묶은 품목의 행 번호를 붙인다."""
    groups, _ = ti.build_ti_groups(items)
    if all("row" in e for entries in groups.values() for e in entries):
        return groups
    rows = defaultdict(list)
    for rec in items:
        iso = ti.normalize_iso(rec.get("ISO NO"))
        dn = ti.dn_number(rec.get("Nominal Diameter")) if iso else None
        if dn is not None:
            rows[(iso, dn)].append(rec["Source Start Row"])
    for key, entries in groups.items():
        if len(entries) != len(rows[key]):
            raise RuntimeError("ti_lst_reader.py 버전이 맞지 않습니다. "
                               "이 프로그램과 같이 받은 ti_lst_reader.py 로 바꿔 주세요.")
        for e, row in zip(entries, rows[key]):
            e["row"] = row
    return groups


class AppendFiles:
    """붙일 품목이 없는 COMMENT 를 적을 TI 파일 찾기: 그 ISO 시트가 있는 파일,
    없으면 같은 구역(ISO 번호 앞부분이 같은) 파일 중 ISO 번호가 가장 가까운 파일."""

    def __init__(self, file_isos):
        self.exact = {}                   # ISO 시트 -> 파일 (파일명 순으로 첫 파일)
        self.by_base = defaultdict(list)  # '...-TKM-ISO' -> [(ISO 번호, 파일)]
        for name in sorted(file_isos):
            for iso in file_isos[name]:
                self.exact.setdefault(iso, name)
                m = ISO_SERIAL_RE.match(iso)
                if m:
                    self.by_base[m.group(1)].append((int(m.group(2)), name))

    def choose(self, iso):
        if iso in self.exact:
            return self.exact[iso]
        m = ISO_SERIAL_RE.match(iso or "")
        candidates = self.by_base.get(m.group(1)) if m else None
        if not candidates:
            return None
        serial = int(m.group(2))
        return min(candidates, key=lambda c: (abs(c[0] - serial), c[1]))[1]


# =========================
# TI-LST 에 적기 (XML 에서 필요한 곳만 고침 - 다른 내용은 그대로)
# =========================
HEADER_ROWS = 3  # TI-LST 머리글 줄 (1~3행) - 필터 단추는 마지막 머리글 줄에 생김
# 시트 XML 에서 autoFilter 뒤에 와야 하는 요소 (이 중 처음 나오는 요소 앞에 autoFilter 를 넣음)
AFTER_AUTOFILTER = ("sortState", "dataConsolidate", "customSheetViews", "mergeCells", "phoneticPr",
                    "conditionalFormatting", "dataValidations", "hyperlinks", "printOptions", "pageMargins",
                    "pageSetup", "headerFooter", "rowBreaks", "colBreaks", "customProperties", "cellWatches",
                    "ignoredErrors", "smartTags", "drawing", "legacyDrawing", "legacyDrawingHF", "drawingHF",
                    "picture", "oleObjects", "controls", "webPublishItems", "tableParts", "extLst")
# workbook.xml 에서 definedNames 뒤에 와야 하는 요소
AFTER_DEFINED_NAMES = ("calcPr", "oleSize", "customWorkbookViews", "pivotCaches", "smartTagPr", "smartTagTypes",
                       "webPublishing", "fileRecoveryPr", "webPublishObjects", "extLst")
THICKNESS_RE = re.compile(r"(\d+(?:[.,]\d+)?)\s*mm(?![a-z])|(?:thk|thickness|두께)\D{0,20}?(\d+(?:[.,]\d+)?)",
                          re.IGNORECASE)


def requested_thickness(comment):
    """COMMENT 에 적힌 두께 ('thickness shall be 50mm' -> 50.0). 없거나 서로 다른 값이 여럿이면 None"""
    values = {ti.to_number(a or b) for a, b in THICKNESS_RE.findall(comment)} - {None}
    return round(values.pop(), 1) if len(values) == 1 else None


def get_attr(attrs, name):
    m = re.search(rf'(?<![\w:.-]){name}\s*=\s*(["\'])(.*?)\1', attrs, re.S)
    return m.group(2) if m else None


def set_attr(attrs, name, value):
    m = re.search(rf'(?<![\w:.-]){name}\s*=\s*(["\'])(.*?)\1', attrs, re.S)
    if m:
        return f"{attrs[:m.start(2)]}{value}{attrs[m.end(2):]}"
    return f'{attrs} {name}="{value}"'


def del_attr(attrs, name):
    return re.sub(rf'\s+{name}\s*=\s*(["\']).*?\1', "", attrs, flags=re.S)


def insert_before_first(text, prefix, names, closing, new):
    """names 중 처음 나오는 요소 앞(없으면 closing 앞)에 new 를 넣는다 (XML 요소 순서를 지키려고)"""
    m = re.search(rf"<{re.escape(prefix)}(?:{'|'.join(names)})\b", text)
    at = m.start() if m else text.rindex(closing)
    return text[:at] + new + text[at:]


def read_rels(zf, part):
    """part 의 관계 파일 -> {Id: (Type, 대상 경로)}"""
    folder, name = posixpath.split(part)
    root = ET.fromstring(zf.read(posixpath.join(folder, "_rels", name + ".rels")))
    rels = {}
    for rel in root:
        if rel.get("TargetMode") == "External":
            continue
        target = rel.get("Target", "")
        path = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(folder, target))
        rels[rel.get("Id")] = (rel.get("Type", ""), path)
    return rels


def find_parts(zf, sheet_name):
    """{'workbook', 'styles', 'sheet': XML 경로, 'index': 시트 순번(0부터)}"""
    workbook = next(path for kind, path in read_rels(zf, "").values() if kind.endswith("/officeDocument"))
    rels = read_rels(zf, workbook)
    styles = next((path for kind, path in rels.values() if kind.endswith("/styles")), None)
    sheets = [s for s in ET.fromstring(zf.read(workbook)).iter() if s.tag.split("}")[-1] == "sheet"]
    for index, sheet in enumerate(sheets):
        if sheet.get("name") == sheet_name:
            rel_id = next(v for k, v in sheet.attrib.items() if k.endswith("}id"))
            return {"workbook": workbook, "styles": styles, "sheet": rels[rel_id][1], "index": index}
    raise ValueError(f"시트를 찾지 못했습니다: {sheet_name}")


def decode_xml(data):
    m = re.match(rb'\s*<\?xml[^>]*encoding=["\']([\w-]+)["\']', data)
    encoding = m.group(1).decode() if m else "utf-8"
    if encoding.lower() in ("utf-8", "utf8") and data.startswith(b"\xef\xbb\xbf"):
        encoding = "utf-8-sig"
    return data.decode(encoding), encoding


class SheetXml:
    """시트 XML 을 정규식으로 필요한 곳만 고친다 (칸·서식, 숨김 행, 필터, 병합, 열 너비, 보기)."""

    def __init__(self, sheet_xml):
        m = re.search(r"<(\w+:)?sheetData\b[^>]*?(/?)>", sheet_xml)
        if not m:
            raise ValueError("sheetData 가 없습니다")
        self.prefix = m.group(1) or ""
        self.p = p = re.escape(self.prefix)
        self.row_re = re.compile(rf"<{p}row\b([^>]*?)(/>|>(.*?)</{p}row>)", re.S)
        self.cell_re = re.compile(rf"<{p}c\b([^>]*?)(/>|>.*?</{p}c>)", re.S)
        self.value_re = re.compile(rf"<{p}(?:v|is|f)\b")
        if m.group(2):  # <sheetData/>
            self.head, self.body, self.tail = sheet_xml[:m.start()], "", sheet_xml[m.end():]
            self.open_tag, self.close_tag = f"<{self.prefix}sheetData>", f"</{self.prefix}sheetData>"
        else:
            end = sheet_xml.index(f"</{self.prefix}sheetData>", m.end())
            self.head, self.body, self.tail = sheet_xml[:m.end()], sheet_xml[m.end():end], sheet_xml[end:]
            self.open_tag = self.close_tag = ""

    def rows(self):
        """(행 번호, 매치) - r 이 없는 행은 앞 행 + 1"""
        current = 0
        for rm in self.row_re.finditer(self.body):
            r = get_attr(rm.group(1), "r")
            current = int(r) if r else current + 1
            yield current, rm

    def cells(self, body):
        """(열 번호, 매치) - r 이 없는 칸은 앞 칸 + 1"""
        current = 0
        for cm in self.cell_re.finditer(body):
            ref = re.match(r"([A-Za-z]+)\d+$", get_attr(cm.group(1), "r") or "")
            current = column_index_from_string(ref.group(1).upper()) if ref else current + 1
            yield current, cm

    def scan(self, col, check_rows):
        """(값이 들어 있는 마지막 행 번호(모든 열 기준), check_rows 중 col 칸에 이미 값이 있는 행)"""
        last, filled = 0, set()
        for row_no, rm in self.rows():
            body = rm.group(3)
            if not body or not self.value_re.search(body):
                continue
            last = row_no
            if row_no in check_rows:
                for c, cm in self.cells(body):
                    if c >= col:
                        if c == col and self.value_re.search(cm.group(0)):
                            filled.add(row_no)
                        break
        return last, filled

    def col_styles(self, max_col):
        """열 기본 서식 {열: 서식 번호} (<cols> 의 style)"""
        styles = {}
        m = re.search(rf"<{self.p}cols\b[^>]*>(.*?)</{self.p}cols>", self.head, re.S)
        for cm in re.finditer(rf"<{self.p}col\b([^>]*?)/?>", m.group(1) if m else ""):
            lo, hi, style = (get_attr(cm.group(1), k) for k in ("min", "max", "style"))
            if lo and hi and style:
                for c in range(int(lo), min(int(hi), max_col) + 1):
                    styles[c] = int(style)
        return styles

    def cell_styles(self, rows, max_col):
        """{행: ({열: 서식 번호}, {값이 있는 열})} - 칸이 없으면 행/열 기본 서식"""
        col_default = self.col_styles(max_col)
        info = {}
        for row_no, rm in self.rows():
            if row_no not in rows:
                continue
            attrs = rm.group(1)
            row_style = get_attr(attrs, "s") if get_attr(attrs, "customFormat") in ("1", "true") else None
            styles = {c: int(row_style) if row_style else col_default.get(c, 0) for c in range(1, max_col + 1)}
            filled = set()
            for c, cm in self.cells(rm.group(3) or ""):
                if c > max_col:
                    break
                styles[c] = int(get_attr(cm.group(1), "s") or 0)
                if self.value_re.search(cm.group(0)):
                    filled.add(c)
            info[row_no] = (styles, filled)
        for row_no in set(rows) - set(info):  # 원본에 없는 행
            info[row_no] = ({c: col_default.get(c, 0) for c in range(1, max_col + 1)}, set())
        return info

    def new_cell(self, col, row, style, text=None):
        p = self.prefix
        ref = f"{get_column_letter(col)}{row}"
        s = f' s="{style}"' if style else ""
        if text is None:
            return f'<{p}c r="{ref}"{s}/>'
        text = ti.ILLEGAL_XML_CHARS_RE.sub("", text)[:ti.EXCEL_MAX_CELL_CHARS]
        return (f'<{p}c r="{ref}"{s} t="inlineStr"><{p}is>'
                f'<{p}t xml:space="preserve">{ti.xml_escape(text)}</{p}t></{p}is></{p}c>')

    def edit_cells(self, body, row, edits):
        """edits {열: (서식, 글자 또는 None)}: 있는 칸은 서식(글자가 있으면 내용도)을 바꾸고,
        없는 칸은 열 순서대로 넣음"""
        todo = dict(edits)
        out, pos, last_end = [], 0, 0
        for c, cm in self.cells(body):
            before = sorted(k for k in todo if k < c)
            if before:
                out += [body[pos:cm.start()]] + [self.new_cell(k, row, *todo.pop(k)) for k in before]
                pos = cm.start()
            if c in todo:
                style, text = todo.pop(c)
                if text is None:  # 내용은 두고 서식만 바꿈
                    start_len = len(f"<{self.prefix}c") + len(cm.group(1))
                    new = f"<{self.prefix}c{set_attr(cm.group(1), 's', style)}{cm.group(0)[start_len:]}"
                else:
                    new = self.new_cell(c, row, style, text)
                out += [body[pos:cm.start()], new]
                pos = cm.end()
            last_end = cm.end()
        if todo:  # 마지막 칸 뒤 (extLst 가 있으면 그 앞)
            at = max(pos, last_end)
            out += [body[pos:at]] + [self.new_cell(k, row, *todo[k]) for k in sorted(todo)]
            pos = at
        out.append(body[pos:])
        return "".join(out)

    def edit_row(self, rm, row, edits, hidden):
        attrs, body = rm.group(1), rm.group(3) or ""
        if edits:
            body = self.edit_cells(body, row, edits)
            spans = get_attr(attrs, "spans")
            if spans:  # 선택 속성: 새 칸을 포함하지 않으면 지움
                try:
                    ranges = [tuple(map(int, s.split(":"))) for s in spans.split()]
                    covered = all(any(a <= c <= b for a, b in ranges) for c in edits)
                except ValueError:
                    covered = False
                if not covered:
                    attrs = del_attr(attrs, "spans")
        if hidden is True:
            attrs = set_attr(attrs, "hidden", "1")
        elif hidden is False:
            attrs = del_attr(attrs, "hidden")
        return f"<{self.prefix}row{attrs}>{body}</{self.prefix}row>" if body else f"<{self.prefix}row{attrs}/>"

    def rewrite(self, edits, hide=None):
        """edits {행: {열: (서식, 글자)}} 로 칸을 바꾸고, hide=(첫 행, 끝 행, 보일 행) 이면 범위 안의 다른 행은
        숨긴다 (필터를 건 결과처럼). 원본에 없는 행은 필요할 때만 만든다."""
        first, last, visible = hide or (0, -1, set())

        def hidden(n):  # True 숨김 / False 보임 / None 그대로
            return (n not in visible) if first <= n <= last else None

        def missing_row(n):
            cells = "".join(self.new_cell(c, n, *edits[n][c]) for c in sorted(edits.get(n, {})))
            state = hidden(n)
            if not cells and not state:
                return ""
            attrs = f' r="{n}"' + (' hidden="1"' if state else "")
            return f"<{self.prefix}row{attrs}>{cells}</{self.prefix}row>" if cells else f"<{self.prefix}row{attrs}/>"

        pieces, pos, prev = [], 0, 0
        for row_no, rm in self.rows():
            gap = "".join(missing_row(n) for n in range(prev + 1, row_no))
            state = hidden(row_no)
            new = rm.group(0)
            if row_no in edits or state is not None:
                new = self.edit_row(rm, row_no, edits.get(row_no), state)
            if gap or new != rm.group(0):
                pieces += [self.body[pos:rm.start()], gap, new]
                pos = rm.end()
            prev = row_no
        pieces.append(self.body[pos:])
        pieces += [missing_row(n) for n in range(prev + 1, max([last, *edits, 0]) + 1)]
        self.body = "".join(pieces)

    def mirror_merges(self, src_col, dst_col, max_row):
        """머리글 줄(max_row 까지)에서 src_col 만 세로로 병합된 범위를 dst_col 에도 똑같이 병합"""
        m = re.search(rf"(<{self.p}mergeCells\b[^>]*?)>(.*?)</{self.p}mergeCells>", self.tail, re.S)
        if not m:
            return
        refs = [get_attr(x.group(1), "ref") for x in re.finditer(rf"<{self.p}mergeCell\b([^>]*?)/?>", m.group(2))]
        bounds = [range_boundaries(ref) for ref in refs if ref]
        letter = get_column_letter(dst_col)
        new = [f'<{self.prefix}mergeCell ref="{letter}{r1}:{letter}{r2}"/>'
               for c1, r1, c2, r2 in bounds
               if c1 == c2 == src_col and r1 < r2 <= max_row
               and not any(a <= dst_col <= c and b <= r2 and r1 <= d for a, b, c, d in bounds)]
        if new:
            start = set_attr(m.group(1), "count", len(bounds) + len(new))
            self.tail = (self.tail[:m.start()] + f"{start}>{m.group(2)}{''.join(new)}</{self.prefix}mergeCells>"
                         + self.tail[m.end():])

    def set_autofilter(self, ref, col_id):
        """ref 범위에 필터를 걸고 col_id 열(0부터)은 '비어 있지 않음'으로 (필터를 건 상태로 저장)"""
        p = self.prefix
        new = (f'<{p}autoFilter ref="{ref}"><{p}filterColumn colId="{col_id}"><{p}customFilters>'
               f'<{p}customFilter operator="notEqual" val=" "/></{p}customFilters></{p}filterColumn></{p}autoFilter>')
        m = re.search(rf"<{self.p}autoFilter\b[^>]*?(?:/>|>.*?</{self.p}autoFilter>)", self.tail, re.S)
        if m:
            self.tail = self.tail[:m.start()] + new + self.tail[m.end():]
        else:
            self.tail = insert_before_first(self.tail, p, AFTER_AUTOFILTER, f"</{p}worksheet>", new)
        sp = re.search(rf"<{self.p}sheetPr\b([^>]*?)(/?>)", self.head)
        if sp:
            self.head = self.head[:sp.start(1)] + set_attr(sp.group(1), "filterMode", "1") + self.head[sp.end(1):]
        else:
            ws = re.search(rf"<{self.p}worksheet\b[^>]*>", self.head)
            self.head = self.head[:ws.end()] + f'<{p}sheetPr filterMode="1"/>' + self.head[ws.end():]

    def set_column_width(self, col, width):
        """col 열 너비를 width 이상으로 (같은 <col> 범위의 다른 열은 나눠서 그대로 두고, 숨김이면 보이게)"""
        p = self.prefix
        own = f'<{p}col min="{col}" max="{col}" width="{width}" customWidth="1"/>'
        m = re.search(rf"<{self.p}cols\b[^>]*>(.*?)</{self.p}cols>", self.head, re.S)
        if not m:  # <cols> 가 없으면 sheetData 바로 앞에 만듦
            sd = re.search(rf"<{self.p}sheetData\b", self.head)
            at = sd.start() if sd else len(self.head)
            self.head = self.head[:at] + f"<{p}cols>{own}</{p}cols>" + self.head[at:]
            return
        out, placed = [], False
        for cm in re.finditer(rf"<{self.p}col\b([^>]*?)/?>", m.group(1)):
            attrs = cm.group(1)
            lo, hi = int(get_attr(attrs, "min") or 0), int(get_attr(attrs, "max") or 0)
            if hi < col or lo > col:
                if lo > col and not placed:
                    out.append(own)
                    placed = True
                out.append(cm.group(0))
                continue
            if lo < col:
                out.append(f"<{p}col{set_attr(set_attr(attrs, 'min', lo), 'max', col - 1)}/>")
            mine = set_attr(set_attr(attrs, "min", col), "max", col)
            current = ti.to_number(get_attr(attrs, "width"))
            if current is None or current < width:  # 원래 더 넓으면 그대로
                mine = set_attr(set_attr(mine, "width", width), "customWidth", "1")
            out.append(f"<{p}col{del_attr(mine, 'hidden')}/>")
            placed = True
            if hi > col:
                out.append(f"<{p}col{set_attr(set_attr(attrs, 'min', col + 1), 'max', hi)}/>")
        if not placed:
            out.append(own)
        self.head = self.head[:m.start(1)] + "".join(out) + self.head[m.end(1):]

    def normal_view(self):
        """'기본' 보기로 (페이지 나누기 미리 보기/페이지 레이아웃 해제, 배율은 기본 보기 배율)"""
        def fix(m):
            attrs = m.group(1)
            if get_attr(attrs, "view") in (None, "normal"):
                return m.group(0)
            attrs = del_attr(attrs, "view")
            zoom = get_attr(attrs, "zoomScaleNormal")
            attrs = set_attr(attrs, "zoomScale", zoom) if zoom else del_attr(attrs, "zoomScale")
            return f"<{self.prefix}sheetView{attrs}{m.group(2)}"
        self.head = re.sub(rf"<{self.p}sheetView\b([^>]*?)(/?>)", fix, self.head)

    def widen_dimension(self, col, row):
        m = re.search(rf'<{self.p}dimension\b[^>]*?\bref\s*=\s*(["\'])([^"\']+)\1', self.head)
        if not m or ":" not in m.group(2):
            return
        start, end = m.group(2).split(":", 1)
        em = re.fullmatch(r"([A-Za-z]+)(\d+)", end)
        if not em:
            return
        new_end = (f"{get_column_letter(max(column_index_from_string(em.group(1).upper()), col))}"
                   f"{max(int(em.group(2)), row)}")
        self.head = self.head[:m.start(2)] + f"{start}:{new_end}" + self.head[m.end(2):]

    def text(self):
        return self.head + self.open_tag + self.body + self.close_tag + self.tail


class StylesXml:
    """styles.xml 에 음영 채우기와, 그 채우기만 바꾼 셀 서식(원래 서식 복사)을 덧붙인다."""

    def __init__(self, styles_xml):
        self.xml = styles_xml
        m = re.search(r"<(\w+:)?styleSheet\b", styles_xml)
        self.prefix = (m.group(1) or "") if m else ""
        p = re.escape(self.prefix)
        self.fills_re = re.compile(rf"(<{p}fills\b[^>]*?)(?:/>|>(.*?)</{p}fills>)", re.S)
        self.xfs_re = re.compile(rf"(<{p}cellXfs\b[^>]*?)(?:/>|>(.*?)</{p}cellXfs>)", re.S)
        fills, xfs = self.fills_re.search(styles_xml), self.xfs_re.search(styles_xml)
        self.ok = bool(fills and xfs)
        self.fill_id = len(re.findall(rf"<{p}fill\b", (fills.group(2) or "") if fills else ""))
        self.xfs = re.findall(rf"<{p}xf\b[^>]*?(?:/>|>.*?</{p}xf>)", (xfs.group(2) or "") if xfs else "", re.S)
        self.new_xfs, self.cache = [], {}

    def highlighted(self, style):
        """style 서식에서 채우기만 음영으로 바꾼 서식 번호"""
        if not self.ok:
            return style
        if style not in self.cache:
            base = self.xfs[style] if style < len(self.xfs) else self.xfs[0]
            m = re.match(rf"<{re.escape(self.prefix)}xf\b([^>]*?)(/?>)", base)
            attrs = set_attr(set_attr(m.group(1), "fillId", self.fill_id), "applyFill", "1")
            self.new_xfs.append(f"<{self.prefix}xf{attrs}{m.group(2)}{base[m.end():]}")
            self.cache[style] = len(self.xfs) + len(self.new_xfs) - 1
        return self.cache[style]

    def text(self):
        p = self.prefix
        fill = (f'<{p}fill><{p}patternFill patternType="solid"><{p}fgColor rgb="{HIGHLIGHT_COLOR}"/>'
                f'<{p}bgColor indexed="64"/></{p}patternFill></{p}fill>')
        xml = self.xml
        m = self.fills_re.search(xml)
        xml = (xml[:m.start()] + f"{set_attr(m.group(1), 'count', self.fill_id + 1)}>{m.group(2) or ''}{fill}"
               f"</{p}fills>" + xml[m.end():])
        m = self.xfs_re.search(xml)
        xml = (xml[:m.start()] + f"{set_attr(m.group(1), 'count', len(self.xfs) + len(self.new_xfs))}>"
               f"{m.group(2) or ''}{''.join(self.new_xfs)}</{p}cellXfs>" + xml[m.end():])
        return xml


def set_filter_database(workbook_xml, sheet_index, sheet_name, ref):
    """workbook.xml 에 시트의 필터 범위 이름(_xlnm._FilterDatabase)을 넣는다 (엑셀/리브레오피스가 필터를 알아보게)"""
    m = re.search(r"<(\w+:)?workbook\b", workbook_xml)
    prefix = (m.group(1) or "") if m else ""
    p = re.escape(prefix)
    quoted = "'" + sheet_name.replace("'", "''") + "'"
    new = (f'<{prefix}definedName name="_xlnm._FilterDatabase" localSheetId="{sheet_index}" hidden="1">'
           f"{ti.xml_escape(quoted)}!{ref}</{prefix}definedName>")
    for dm in re.finditer(rf"<{p}definedName\b([^>]*)>.*?</{p}definedName>", workbook_xml, re.S):
        if (get_attr(dm.group(1), "name") == "_xlnm._FilterDatabase"
                and get_attr(dm.group(1), "localSheetId") == str(sheet_index)):
            return workbook_xml[:dm.start()] + new + workbook_xml[dm.end():]
    empty = re.search(rf"<{p}definedNames\s*/>", workbook_xml)
    if empty:
        return workbook_xml[:empty.start()] + f"<{prefix}definedNames>{new}</{prefix}definedNames>" + \
            workbook_xml[empty.end():]
    close = f"</{prefix}definedNames>"
    if close in workbook_xml:
        at = workbook_xml.index(close)
        return workbook_xml[:at] + new + workbook_xml[at:]
    return insert_before_first(workbook_xml, prefix, AFTER_DEFINED_NAMES, f"</{prefix}workbook>",
                               f"<{prefix}definedNames>{new}</{prefix}definedNames>")


def existing_values(path, sheet_name, col):
    """원본 TI-LST 의 COMMENT 열에 이미 있는 값 {행: 값}"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet_name]
        ws.reset_dimensions()
        return {r: v for r, (v,) in enumerate(ws.iter_rows(min_col=col, max_col=col, values_only=True), start=1)
                if v not in (None, "")}
    finally:
        wb.close()


def write_ti_comments(src, sheet_name, out_path, row_texts, append_texts):
    """원본 TI-LST 를 복사하며 COMMENT 를 넣는다. 머리글 'Comment', COMMENT 행 음영, COMMENT 행만 보이는 필터,
    기본 보기까지 설정하고, 바뀐 XML(시트·서식·통합 문서)만 바꿔 저장한다 (그림·머리글/바닥글 등은 그대로).

    row_texts: {품목 행: [COMMENT...]} / append_texts: 데이터 끝 다음 행부터 차례로 적을 글자
    돌려주는 값: ({품목 행: 기존 값 있었는지}, [추가한 행 번호])"""
    col = column_index_from_string(COMMENT_COLUMN)
    letter = get_column_letter(col)
    with zipfile.ZipFile(src) as zin:
        parts = find_parts(zin, sheet_name)
        sheet_xml, sheet_encoding = decode_xml(zin.read(parts["sheet"]))
        sheet = SheetXml(sheet_xml)
        last_row, filled = sheet.scan(col, set(row_texts))

        # 품목 행의 COMMENT 열에 원래 값이 있으면 그 뒤에 붙임 (값이 있는 행이 있을 때만 다시 읽음)
        old = existing_values(src, sheet_name, col) if filled else {}
        values, had_old = {}, {}
        for row, texts in row_texts.items():
            text = " / ".join(dict.fromkeys(texts))
            before = comment_text(old.get(row))
            had_old[row] = bool(before) and text not in before
            values[row] = f"{before} / {text}" if had_old[row] else (before or text)
        first = max(last_row, max(values, default=0)) + 1  # 원본 데이터가 끝난 다음 행
        appended = list(range(first, first + len(append_texts)))
        values.update(zip(appended, append_texts))
        last = max(last_row, max(values, default=0))

        # COMMENT 칸 + 행 음영 (원래 서식에 채우기만 바꿈)
        header_rows = range(1, HEADER_ROWS + 1)
        cells = sheet.cell_styles(set(values) | (set(header_rows) if COMMENT_HEADER else set()), col)
        styles = None
        if HIGHLIGHT_COLOR and parts["styles"]:
            styles_xml, styles_encoding = decode_xml(zin.read(parts["styles"]))
            styles = StylesXml(styles_xml)
        edits = defaultdict(dict)
        for row, text in values.items():
            row_styles = cells[row][0]
            for c in (range(1, col + 1) if styles else [col]):
                style = styles.highlighted(row_styles[c]) if styles else row_styles[c]
                edits[row][c] = (style, text if c == col else None)
        # 머리글: N열(Note)과 같은 서식, N열 이름이 있는 줄에 'Comment' (N열이 병합돼 있으면 똑같이 병합)
        if COMMENT_HEADER:
            note = col - 1
            name_row = next((r for r in header_rows if note in cells[r][1]), HEADER_ROWS)
            for r in header_rows:
                edits[r][col] = (cells[r][0][note], COMMENT_HEADER if r == name_row else None)
            sheet.mirror_merges(note, col, HEADER_ROWS)

        sheet.rewrite(edits, (HEADER_ROWS + 1, last, set(values)) if FILTER_COMMENT_ROWS else None)
        sheet.widen_dimension(col, last)
        if FILTER_COMMENT_ROWS:
            sheet.set_autofilter(f"A{HEADER_ROWS}:{letter}{last}", col - 1)
        if COMMENT_COLUMN_WIDTH:
            sheet.set_column_width(col, COMMENT_COLUMN_WIDTH)
        if NORMAL_VIEW:
            sheet.normal_view()

        replaced = {parts["sheet"]: sheet.text().encode(sheet_encoding)}
        if styles and styles.new_xfs:
            replaced[parts["styles"]] = styles.text().encode(styles_encoding)
        if FILTER_COMMENT_ROWS:
            workbook_xml, workbook_encoding = decode_xml(zin.read(parts["workbook"]))
            workbook_xml = set_filter_database(workbook_xml, parts["index"], sheet_name,
                                               f"$A${HEADER_ROWS}:${letter}${last}")
            replaced[parts["workbook"]] = workbook_xml.encode(workbook_encoding)
        for data in replaced.values():
            ParserCreate().Parse(data, True)  # 고친 XML 이 깨지지 않았는지 확인 (깨졌으면 오류 -> 저장 안 함)

        # 나머지 파일(그림, 공유 문자열 등)은 원본 그대로 복사
        tmp_path = out_path.with_name(out_path.name + ".tmp")
        try:
            with zipfile.ZipFile(tmp_path, "w") as zout:
                for info in zin.infolist():
                    copy = zipfile.ZipInfo(info.filename, info.date_time)
                    copy.compress_type = info.compress_type
                    copy.external_attr = info.external_attr
                    zout.writestr(copy, replaced[info.filename] if info.filename in replaced
                                  else zin.read(info.filename))
            os.replace(tmp_path, out_path)
        finally:
            if tmp_path.exists():
                tmp_path.unlink()
    return had_old, appended


# =========================
# 실행
# =========================
def is_eht_list(path):
    """ti_lst_reader 처럼 파일명에 EHT 가 있고 EHT LIST 머리글이 있으면 TI-LST 가 아님"""
    if not ti.EHT_FILE_KEYWORD or ti.EHT_FILE_KEYWORD.lower() not in path.name.lower():
        return False
    try:
        return ti.read_eht_list(path) is not None
    except Exception:
        return False


def find_ti_files(base_dir, out_dir):
    """ti_lst_reader 와 같은 규칙으로 TI-LST 파일 찾기 (결과 폴더, COMMENT/EHT LIST 파일 제외)"""
    files = base_dir.rglob("*.xlsx") if ti.RECURSIVE else base_dir.glob("*.xlsx")
    skip = [k.lower() for k in [*ti.SKIP_KEYWORDS, COMMENT_FILE_KEYWORD]]
    return sorted(f for f in files
                  if not any(k in f.name.lower() for k in skip) and out_dir not in f.parents and not is_eht_list(f))


def parse_ti_files(paths):
    """TI-LST 읽기 (ti_lst_reader 와 같은 방식, 병렬) -> 품목 목록"""
    items = []
    for file_items, _, summary, _ in ti.process_files(paths):
        if summary["Status"] == "ERROR":
            continue
        items.extend(file_items)
    return items


def item_info(item):
    """반영 결과에 남길 TI 품목 정보"""
    return {"TI 파일": item["Source File"], "TI 행": item["Source Start Row"], "품명": item["Item Name"],
            "Size (mm)": item["Size (mm)"], "Nominal Diameter": item["Nominal Diameter"],
            "두께 (mm)": item["Insulation Thickness (mm)"], "LINE NO": item["LINE NO"]}


def log_row(result, rec, **extra):
    row = {
        "결과": result, "COMMENT": rec["COMMENT"], "ISO NO": rec.get("ISO NO"), "DN": rec.get("DN"),
        "판정": rec.get("판정"), "EHT 라인": rec.get("EHT 라인"),
        "COMMENT 파일": rec["COMMENT 파일"], "COMMENT 행": rec["COMMENT 행"],
    }
    row.update(extra)
    return row


def main():
    print("=== COMMENT -> TI-LST ===")
    started = time.perf_counter()
    base_dir = Path(__file__).resolve().parent
    out_dir = base_dir / OUTPUT_FOLDER
    print(f"작업 폴더: {base_dir}")

    # 1) COMMENT 파일 읽기 (끌어다 놓은 파일, 없으면 폴더에서 파일명에 COMMENT 가 있는 엑셀)
    comment_files = [Path(a) for a in sys.argv[1:] if a.lower().endswith(".xlsx")] or sorted(
        f for f in base_dir.glob("*.xlsx")
        if COMMENT_FILE_KEYWORD.lower() in f.name.lower() and not f.name.startswith("~$"))
    comments = []
    for path in comment_files:
        found = read_comment_file(path)
        if found is None:
            print(f"!! EHT비교 시트를 찾지 못함: {path.name}")
            continue
        records, sheet, col = found
        print(f"COMMENT 파일: {path.name} ('{sheet}' 시트 {col}열, COMMENT {len(records)}건)")
        comments.extend(records)
    if not comments:
        print(f"COMMENT 가 없습니다. 파일명에 '{COMMENT_FILE_KEYWORD}' 가 들어간 EHT비교 파일을 이 폴더에 넣어 주세요.")
        input("엔터를 누르면 종료합니다...")
        return

    # 2) TI-LST 읽기: COMMENT 에 나온 파일 먼저, 붙일 품목이 없는 COMMENT 가 있으면 나머지도 (같은 구역 파일 찾기)
    out_dir.mkdir(exist_ok=True)
    ti_files = find_ti_files(base_dir, out_dir)
    by_name = {}
    for f in ti_files:
        by_name.setdefault(f.name, f)
    named = {name for rec in comments for name in split_names(rec.get("TI 파일"), by_name)}
    first = [by_name[name] for name in sorted(named) if name in by_name]
    print(f"TI-LST 읽는 중... (COMMENT 에 나온 파일 {len(first)}개)")
    items = parse_ti_files(first) if first else []

    groups = ti_groups(items)
    targets = [comment_targets(rec, groups, by_name) for rec in comments]
    if any(not picked and not split_names(rec.get("TI 파일")) for rec, picked in zip(comments, targets)):
        rest = [f for f in ti_files if f.name not in named]
        if rest:
            print(f"TI 파일이 없는 COMMENT(TI에 ISO 없음 등)가 있어 같은 구역 TI-LST 를 찾으려고 나머지도 읽습니다 "
                  f"({len(rest)}개)")
            items += parse_ti_files(rest)

    item_of = {(rec["Source File"], rec["Source Start Row"]): rec for rec in items}
    path_of = {rec["Source File"]: Path(rec["Source Path"]) for rec in items}
    sheet_of = {rec["Source File"]: rec["Sheet Name"] for rec in items}
    file_isos = defaultdict(set)
    for rec in items:
        if rec["ISO NO"]:
            file_isos[rec["Source File"]].add(ti.normalize_iso(rec["ISO NO"]))
    append_files = AppendFiles(file_isos)

    # 3) 파일별로 적을 내용 모으기
    row_texts = defaultdict(lambda: defaultdict(list))   # 파일 -> 품목 행 -> [COMMENT]
    row_comments = defaultdict(list)                      # (파일, 행) -> [COMMENT 행 레코드]
    appends = defaultdict(list)                           # 파일 -> [(글자, COMMENT 행 레코드)]
    log = []
    for rec, picked in zip(comments, targets):
        if picked:
            # COMMENT 에 적힌 두께와 이미 같은 두께인 품목은 빼고 적음 (예: 50mm 로 바꾸라는 COMMENT -> 50mm 품목 제외)
            thk = requested_thickness(rec["COMMENT"]) if SKIP_SAME_THICKNESS else None
            for e in picked:
                num = ti.to_number(e["thk"])  # '30+30' 처럼 숫자 하나가 아닌 두께는 None (그 품목에는 적음)
                if thk is not None and num is not None and round(num, 1) == thk:
                    log.append(log_row(f"적지 않음 (이미 {ti.tidy_number(thk)}mm)", rec,
                                       **item_info(item_of[(e["file"], e["row"])])))
                    continue
                row_texts[e["file"]][e["row"]].append(rec["COMMENT"])
                row_comments[(e["file"], e["row"])].append(rec)
            continue
        # 붙일 품목이 없으면 데이터 끝에 추가: TI 파일이 적혀 있으면 그 파일, 없으면 같은 구역 파일
        names = split_names(rec.get("TI 파일"), by_name)
        if names:
            name = next((n for n in names if n in sheet_of), None)
            if name is None:
                log.append(log_row("TI 파일 없음 (폴더에 없거나 읽지 못함)", rec, **{"TI 파일": ", ".join(names)}))
                continue
        else:
            iso = comment_iso(rec)
            name = append_files.choose(iso) if iso else None
            if name is None:
                log.append(log_row("TI 파일 없음 (같은 구역 TI-LST 를 찾지 못함)", rec))
                continue
        line = ti.clean_text(rec.get("EHT 라인")) or ti.clean_text(rec.get("TI LINE NO"))
        text = APPEND_FORMAT.format(iso=ti.clean_text(rec.get("ISO NO")), line=line,
                                    dn=ti.clean_text(rec.get("DN")), comment=rec["COMMENT"])
        appends[name].append((text, rec))

    # 4) TI-LST 저장 (원본은 그대로, OUTPUT 폴더에 '원래이름_COMMENT.xlsx')
    saved = 0
    for name in sorted(set(row_texts) | set(appends)):
        src = path_of[name]
        out_path = out_dir / f"{src.stem}{OUTPUT_SUFFIX}{src.suffix}"
        try:
            had_old, appended = write_ti_comments(src, sheet_of[name], out_path, row_texts.get(name, {}),
                                                  [text for text, _ in appends.get(name, [])])
        except Exception as e:
            print(f"!! 저장 실패: {name} ({e})")
            for row in row_texts.get(name, {}):
                for rec in row_comments[(name, row)]:
                    log.append(log_row(f"저장 실패: {e}", rec, **{"TI 파일": name, "TI 행": row}))
            for _, rec in appends.get(name, []):
                log.append(log_row(f"저장 실패: {e}", rec, **{"TI 파일": name}))
            continue
        saved += 1
        print(f"저장: {out_path.name} (품목 {len(row_texts.get(name, {}))}행, 데이터 끝에 추가 {len(appended)}행)")
        for row in sorted(row_texts.get(name, {})):
            result = "품목 행에 적음 (기존 값 뒤에 붙임)" if had_old[row] else "품목 행에 적음"
            for rec in row_comments[(name, row)]:
                log.append(log_row(result, rec, **item_info(item_of[(name, row)]), **{"저장 파일": out_path.name}))
        for row, (_, rec) in zip(appended, appends.get(name, [])):
            log.append(log_row("데이터 끝에 추가", rec, **{
                "TI 파일": name, "TI 행": row, "저장 파일": out_path.name}))

    # 5) 어디에 적었는지 기록
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    log_path = out_dir / f"COMMENT_반영결과_{timestamp}.xlsx"
    columns = ti.order_columns(log, LOG_COLUMNS)
    ti.write_xlsx(log_path, [("반영결과", columns, [[rec.get(c) for c in columns] for rec in log])])

    print("\n=== 작업 완료 ===")
    print(f"COMMENT {len(comments)}건 -> TI-LST {saved}개 저장 ({out_dir})")
    skipped = sum(1 for rec in log if rec["결과"].startswith("적지 않음"))
    if skipped:
        print(f"이미 COMMENT 의 두께인 품목 {skipped}행은 적지 않음 - {log_path.name} 의 '결과' 열 확인")
    failed = sum(1 for rec in log if not rec["결과"].startswith(("품목 행", "데이터 끝", "적지 않음")))
    if failed:
        print(f"!! 적지 못한 COMMENT {failed}건 - {log_path.name} 의 '결과' 열 확인")
    print(f"반영 결과: {log_path}")
    print(f"소요 시간: {time.perf_counter() - started:.1f}초")
    input("엔터를 누르면 종료합니다...")


if __name__ == "__main__":
    freeze_support()  # exe로 묶어 실행할 때 병렬 처리용
    try:
        main()
    except Exception as e:
        print("### 프로그램 치명적 오류 ###")
        print(e)
        print(traceback.format_exc())
        input("엔터를 누르면 종료합니다...")
