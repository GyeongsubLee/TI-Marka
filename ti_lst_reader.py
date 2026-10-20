from pathlib import Path
import math
import os
import re
import time
import traceback
import unicodedata
import zipfile
from collections import Counter, defaultdict
from concurrent.futures import ProcessPoolExecutor
from datetime import datetime
from itertools import islice, zip_longest
from multiprocessing import freeze_support
import warnings

# 경고 숨김 (데이터 처리에는 영향 없는 openpyxl 경고)
warnings.filterwarnings("ignore", message=r"Cannot parse header or footer.*")
warnings.filterwarnings("ignore", message=r"DrawingML support is incomplete.*")

import openpyxl
from openpyxl.utils import get_column_letter


# =========================
# 사용자 설정
# =========================
RECURSIVE = False   # True: 하위 폴더까지 검색 / False: 현재 폴더만 검색
SKIP_KEYWORDS = ["가공본", "통합정리", "~$", "COMMENT"]  # 결과물/임시파일/COMMENT 파일 제외
DATA_START_ROW = 5   # 실제 데이터 시작 추정 행
VALIDATION_CHECK_MAX_COL = 14  # A:N 기준 검토 (이 열까지만 읽음)
MAX_WORKERS = 0  # 동시에 처리할 프로세스 수 (0: 자동 = CPU 코어 수, 최대 8 / 1: 병렬 처리 안 함)
EHT_FILE_KEYWORD = "EHT"  # 파일명에 이 글자가 있는 EHT LIST 를 같이 넣으면 TI 와 비교한 시트(EHT비교)를 추가
EHT_LENGTH_TOLERANCE = 0.1  # EHT 길이와 TI 파이프 길이 차이가 이 값(m)보다 크면 '길이 차이'


# =========================
# Pipe OD -> Nominal Diameter (DN) 변환표
# =========================
OD_TO_DN_MAP = {
    10.29: "DN 6", 13.72: "DN 8", 17.14: "DN 10",
    18.00: "DN 15", 21.34: "DN 15",
    25.00: "DN 20", 26.67: "DN 20",
    32.00: "DN 25", 33.40: "DN 25",
    38.00: "DN 32", 42.16: "DN 32",
    45.00: "DN 40", 48.26: "DN 40",
    57.00: "DN 50", 60.32: "DN 50",
    73.02: "DN 65", 76.00: "DN 65",
    88.90: "DN 80", 89.00: "DN 80",
    101.60: "DN 90",
    108.00: "DN 100", 114.30: "DN 100",
    141.30: "DN 125",
    159.00: "DN 150", 168.28: "DN 150",
    219.00: "DN 200", 219.08: "DN 200",
    273.00: "DN 250",
    323.80: "DN 300", 325.00: "DN 300",
    355.60: "DN 350", 377.00: "DN 350",
    406.40: "DN 400", 426.00: "DN 400",
    457.20: "DN 450",
    508.00: "DN 500", 530.00: "DN 500",
    558.80: "DN 550",
    609.60: "DN 600", 630.00: "DN 600",
    660.40: "DN 650",
    711.20: "DN 700",
    762.00: "DN 750",
    813.00: "DN 800", 820.00: "DN 800",
    863.60: "DN 850", 864.00: "DN 850",
    914.00: "DN 900", 914.40: "DN 900",
    965.00: "DN 950",
    1016.00: "DN 1000", 1020.00: "DN 1000",
    1067.00: "DN 1050",
    1117.60: "DN 1100", 1118.00: "DN 1100",
    1168.00: "DN 1150", 1168.40: "DN 1150",
    1219.00: "DN 1200", 1220.00: "DN 1200",
    1270.00: "DN 1250",
    1321.00: "DN 1300",
    1420.00: "DN 1400", 1422.00: "DN 1400", 1422.40: "DN 1400",
    1524.00: "DN 1500",
    1625.60: "DN 1600", 1626.00: "DN 1600",
    1727.00: "DN 1700",
    1829.00: "DN 1800",
    1930.00: "DN 1900", 1930.40: "DN 1900",
    2032.00: "DN 2000",
}

DN_VALUES = set(OD_TO_DN_MAP.values())  # 유효한 DN 호칭 목록 (DN 6 ~ DN 2000)


# =========================
# Tube 규격 -> Nominal Diameter (DN) 변환표
# (품명의 "TUBING SIZE: 3/4 INCH" 인치 값으로 찾음. 표에 없는 규격은 빈칸)
# =========================
TUBE_SIZE_TO_DN = {
    "1/8": "-",
    "1/4": "DN 8",
    "3/8": "DN 10",
    "1/2": "DN 15",
    "3/4": "DN 20",
    "1": "DN 25",
}


def clean_text(x):
    if x is None:
        return ""
    if isinstance(x, str):
        return " ".join(x.replace("\n", " ").replace("\r", " ").split())
    return str(x)


def get_target_sheet(workbook):
    exact_name = "Техномонтажная ведомость"
    if exact_name in workbook.sheetnames:
        return workbook[exact_name]

    for name in workbook.sheetnames:
        lname = name.lower()
        if "ведомость" in lname:
            return workbook[name]

    for name in workbook.sheetnames:
        lname = name.lower()
        if "assembly" in lname or "insulation" in lname:
            return workbook[name]

    return None


def read_sheet_rows(file_path):
    """대상 시트의 A:N 값만 읽어 clean_text 처리된 행 목록으로 돌려준다.

    read_only 모드는 Cover Page/그림/서식을 통째로 메모리에 올리지 않고
    대상 시트 XML만 순서대로 읽기 때문에 일반 모드보다 훨씬 빠르다.
    병합셀은 왼쪽 위 셀에만 값이 있고, 라인/섹션/꼬리말 판정은 모두 A열 기준이라 펼치지 않는다.
    """
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=False)
    try:
        ws = get_target_sheet(wb)
        if ws is None:
            raise ValueError("대상 시트를 찾지 못했습니다. ('Техномонтажная ведомость' 등)")

        sheet_name = ws.title
        sheet_cols = ws.max_column  # 파일에 기록된 사용 범위 (없으면 None)

        # 일부 프로그램이 만든 파일은 사용 범위 정보가 틀려 행이 잘려 읽히므로 끝까지 읽게 함
        ws.reset_dimensions()
        rows = [
            [clean_text(v) for v in row]
            for row in ws.iter_rows(max_col=VALIDATION_CHECK_MAX_COL, values_only=True)
        ]
    finally:
        wb.close()  # read_only 모드는 파일 핸들을 열어두므로 반드시 닫음

    return sheet_name, rows, sheet_cols or VALIDATION_CHECK_MAX_COL


def get_row_values(mat, row_no):
    # read_sheet_rows 에서 이미 A:N 으로 맞추고 clean_text 처리해 둠
    return mat[row_no - 1]


def is_pipeline_header(text):
    t = clean_text(text)
    return "Pipeline / Трубопровод -" in t


def is_footer_row(text):
    t = clean_text(text)
    return t.startswith("Sheet /") or t.startswith("Лист")


def is_meaningful_row(row_vals):
    return any(clean_text(v) for v in row_vals)


def get_sheet_no(line_no):
    """LINE NO 오른쪽 끝의 시트번호 ('12200-24-4312/341-P44-0018-ASB0T00NN51-01 (H)' -> '01').

    끝에 붙은 '(H)', 'H(REGEN)' 같은 꼬리표는 건너뛰고, 오른쪽부터 '-숫자(1~2자리)'로 끝나는 부분을 찾는다.
    TRIM 라인('.../351-F-101-XT-01')처럼 구조가 달라도 끝자리만 보므로 같은 방식으로 찾는다. 없으면 빈칸.
    """
    for token in reversed(clean_text(line_no).split()):
        m = re.search(r"-(\d{1,2})$", re.sub(r"\([^()]*\)", "", token))
        if m:
            return m.group(1)
    return ""


def parse_pipeline_header(text):
    t = clean_text(text)

    marker = "Pipeline / Трубопровод -"
    idx = t.find(marker)
    if idx >= 0:
        t = t[idx:]

    result = {
        "Line Header": t,
        "LINE NO": "",
        "ISO NO": "",
        "INDOOR/OUTDOOR": "",
        "Excluded": "N",
        "Exclusion Letter": "",
    }

    m = re.search(
        r"Pipeline\s*/\s*Трубопровод\s*-\s*(.+?)\s+([A-Z]+)\s*\((INDOOR|OUTDOOR)\)",
        t,
        flags=re.IGNORECASE,
    )
    if m:
        base_line_no = m.group(1).strip()
        symbol = m.group(2).strip().upper()
        inout = m.group(3).strip().upper()
        result["LINE NO"] = f"{base_line_no} ({symbol})"
        result["INDOOR/OUTDOOR"] = inout
    else:
        m2 = re.search(
            r"Pipeline\s*/\s*Трубопровод\s*-\s*(.+?)\s*\((GCC-[^)]+ISO-\d+)\)",
            t,
            flags=re.IGNORECASE,
        )
        if m2:
            result["LINE NO"] = m2.group(1).strip()

    m_iso = re.search(r"(GCC-[A-Z0-9-]*TKM-ISO-\d+)", t, flags=re.IGNORECASE)
    if m_iso:
        result["ISO NO"] = m_iso.group(1)
        # ISO NO 뒤에 시트번호(LINE NO 오른쪽 끝 숫자)를 붙임: ...-TKM-ISO-00094 -> ...-TKM-ISO-00094-01
        sheet_no = get_sheet_no(result["LINE NO"])
        if sheet_no:
            result["ISO NO"] += f"-{sheet_no}"

    if "excluded according to the letter" in t.lower() or "исключено согласно письму" in t.lower():
        result["Excluded"] = "Y"
        matches = re.findall(r"(GCC-[A-Z0-9-]+)", t, flags=re.IGNORECASE)
        candidates = [m for m in matches if "ISO" not in m.upper()]
        if candidates:
            result["Exclusion Letter"] = candidates[-1]

    return result


def parse_size_to_float(size_text):
    s = clean_text(size_text)
    if not s:
        return None

    s = s.replace(" ", "")

    if re.fullmatch(r"\d{1,3}(,\d{3})+", s):
        s = s.replace(",", "")
        try:
            return float(s)
        except Exception:
            return None

    s = s.replace(",", ".")
    try:
        return float(s)
    except Exception:
        return None


def get_nominal_diameter(item_name, od_text, tolerance=0.3):
    name = clean_text(item_name).lower()
    if not name:
        return ""

    # Tube 는 외경(Size) 대신 품명의 "TUBING SIZE: 3/4 INCH" 인치 규격으로 DN 지정
    if "tube" in name:
        m = re.search(r"tubing size:\s*([\d/ -]+?)\s*inch", name)
        return TUBE_SIZE_TO_DN.get(m.group(1), "") if m else ""

    is_valve = "valve" in name
    if not (is_valve or "pipe" in name or "труба" in name):
        return ""

    od_val = parse_size_to_float(od_text)
    if od_val is None:
        return ""

    # Valve 는 Size 칸에 외경이 아니라 호칭경이 적혀 있음 (25,0 -> DN 25).
    # 외경표에는 25.00(DN 20), 32.00(DN 25) 같은 값이 있어 표로 찾으면 틀리므로 먼저 확인한다.
    # DN 값이 아니면(외경을 적은 경우 등) 아래 외경표로 찾는다.
    if is_valve and od_val.is_integer() and f"DN {int(od_val)}" in DN_VALUES:
        return f"DN {int(od_val)}"

    for ref_od, dn in OD_TO_DN_MAP.items():
        if abs(od_val - ref_od) <= tolerance:
            return dn

    return ""


def parse_insulation_workbook(file_path):
    sheet_name, mat, cols = read_sheet_rows(file_path)
    rows = len(mat)

    # ----- row 4가 Section인지, 바로 Pipeline Header인지 동적 판별 -----
    section_value = ""
    row4_a = clean_text(mat[3][0]) if rows >= 4 and cols >= 1 else ""

    if row4_a and not is_pipeline_header(row4_a):
        # 일반 케이스: row4 = Section
        section_value = row4_a
        start_row = 5
        covered_rows = set([1, 2, 3, 4])
    else:
        # 예외 케이스: row4 = 첫 Pipeline Header
        section_value = ""
        start_row = 4
        covered_rows = set([1, 2, 3])

    item_records = []
    excluded_headers = []
    parse_issues = []

    current_header = {
        "Line Header": "",
        "LINE NO": "",
        "ISO NO": "",
        "INDOOR/OUTDOOR": "",
        "Excluded": "",
        "Exclusion Letter": "",
    }

    r = start_row
    while r <= rows:
        row = get_row_values(mat, r)
        a, b, c, d, e, f, g, h, i, j, k, l, m, n = row

        if is_footer_row(a):
            covered_rows.add(r)
            r += 1
            continue

        if not is_meaningful_row(row):
            covered_rows.add(r)
            r += 1
            continue

        if is_pipeline_header(a):
            current_header = parse_pipeline_header(a)
            covered_rows.add(r)

            # 시트번호를 못 찾으면 ISO NO 가 시트번호 없이 남으므로 검토이슈로 남김
            if current_header["LINE NO"] and current_header["ISO NO"] and not get_sheet_no(current_header["LINE NO"]):
                parse_issues.append({
                    "Source File": file_path.name,
                    "Source Path": str(file_path),
                    "Sheet Name": sheet_name,
                    "Issue Type": "SHEET_NO_NOT_FOUND",
                    "Message": "LINE NO 끝에서 시트번호를 찾지 못해 ISO NO에 시트번호를 붙이지 못했습니다.",
                    "Row No": r,
                    "Source Start Row": r,
                    "Source End Row": r,
                    "LINE NO": current_header["LINE NO"],
                    "ISO NO": current_header["ISO NO"],
                    "Item Name": "",
                    "Row Snapshot": a,
                })

            if current_header["Excluded"] == "Y":
                excluded_headers.append({
                    "Source File": file_path.name,
                    "Source Path": str(file_path),
                    "Sheet Name": sheet_name,
                    "Section": section_value,
                    **current_header,
                    "Source Start Row": r,
                    "Source End Row": r,
                })

            r += 1
            continue

        # 품목행
        if b and not b.startswith("Name/") and not b.startswith("Pipeline /"):
            rec = {
                "Source File": file_path.name,
                "Source Path": str(file_path),
                "Sheet Name": sheet_name,
                "Section": section_value,

                "Line Header": current_header.get("Line Header", ""),
                "LINE NO": current_header.get("LINE NO", ""),
                "ISO NO": current_header.get("ISO NO", ""),
                "INDOOR/OUTDOOR": current_header.get("INDOOR/OUTDOOR", ""),
                "Excluded": current_header.get("Excluded", ""),
                "Exclusion Letter": current_header.get("Exclusion Letter", ""),

                "Brand/Position": a,
                "Item Name": b,
                "Size (mm)": c,
                "Nominal Diameter": get_nominal_diameter(b, c),
                "Length/Height (m)": d,
                "Qty": e,
                "Temperature (°C)": f,
                "Purpose/Location": g,
                "Insulation Material": h,
                "Insulation Thickness (mm)": i,
                "Covering Thickness (mm)": j,
                "Surface (m²)": k,
                "Volume of Insulation Layer (m³)": l,
                "Document Designation": m,
                "Note": n,

                "Insulation Detail": "",
                "Insulation Detail Thickness (mm)": "",
                "Covering Material": "",
                "Covering Surface (m²)": "",

                "Source Start Row": r,
                "Source End Row": r,
            }

            rr = r + 1
            while rr <= rows:
                next_row = get_row_values(mat, rr)
                na, nb, nc, nd, ne, nf, ng, nh, ni, nj, nk, nl, nm, nn = next_row

                if is_pipeline_header(na) or nb or is_footer_row(na):
                    break

                # 병합셀을 펼치지 않으므로 A열에만 값이 있는 행(섹션/반복 표머리/메모 등)도
                # 상세행으로 흡수하지 않고 끊는다 → 미분류 행으로 에러로그에 기록됨
                if na:
                    break

                if not is_meaningful_row(next_row):
                    # 완전 공란 행이면 item 범위에 넣지 않고 넘어감
                    break

                if any(next_row):
                    low_h = nh.lower() if nh else ""

                    if nh and ("thickness" in low_h or "толщина" in low_h):
                        if not rec["Insulation Detail"]:
                            rec["Insulation Detail"] = nh
                        if ni and not rec["Insulation Detail Thickness (mm)"]:
                            rec["Insulation Detail Thickness (mm)"] = ni
                        if nl and not rec["Volume of Insulation Layer (m³)"]:
                            rec["Volume of Insulation Layer (m³)"] = nl
                    elif nh:
                        if (nj or nk) and not rec["Covering Material"]:
                            rec["Covering Material"] = nh
                            if nj and not rec["Covering Thickness (mm)"]:
                                rec["Covering Thickness (mm)"] = nj
                            if nk and not rec["Covering Surface (m²)"]:
                                rec["Covering Surface (m²)"] = nk
                        elif not rec["Insulation Detail"]:
                            rec["Insulation Detail"] = nh

                    if ni and not rec["Insulation Detail Thickness (mm)"] and rec["Insulation Detail"]:
                        rec["Insulation Detail Thickness (mm)"] = ni

                rec["Source End Row"] = rr
                rr += 1

            # Source Start/End 범위 전체를 covered 처리
            for rr_cov in range(rec["Source Start Row"], rec["Source End Row"] + 1):
                covered_rows.add(rr_cov)

            # 기본 검토: ITEM인데 LINE/ISO가 비어 있으면 이슈 기록
            if not rec["LINE NO"] or not rec["ISO NO"]:
                parse_issues.append({
                    "Source File": file_path.name,
                    "Source Path": str(file_path),
                    "Sheet Name": sheet_name,
                    "Issue Type": "MISSING_LINE_OR_ISO_ON_ITEM",
                    "Message": "ITEM은 추출되었으나 LINE NO 또는 ISO NO가 비어 있습니다.",
                    "Row No": rec["Source Start Row"],
                    "Source Start Row": rec["Source Start Row"],
                    "Source End Row": rec["Source End Row"],
                    "LINE NO": rec["LINE NO"],
                    "ISO NO": rec["ISO NO"],
                    "Item Name": rec["Item Name"],
                    "Row Snapshot": " | ".join(get_row_values(mat, rec["Source Start Row"])),
                })

            item_records.append(rec)
            r = rr
            continue

        # 여기까지 왔는데 의미 있는 행이면 파서가 놓친 행
        parse_issues.append({
            "Source File": file_path.name,
            "Source Path": str(file_path),
            "Sheet Name": sheet_name,
            "Issue Type": "UNPARSED_ROW_DURING_SCAN",
            "Message": "의미 있는 행이지만 ITEM/HEADER/FOOTER로 분류되지 않았습니다.",
            "Row No": r,
            "Source Start Row": r,
            "Source End Row": r,
            "LINE NO": current_header.get("LINE NO", ""),
            "ISO NO": current_header.get("ISO NO", ""),
            "Item Name": b,
            "Row Snapshot": " | ".join(row),
        })
        r += 1

    # 후검증: Source Start/End 범위를 기준으로 커버되지 않은 의미 있는 행 전부 찾기
    for chk_row in range(start_row, rows + 1):
        row_vals = get_row_values(mat, chk_row)
        a = row_vals[0]

        if chk_row in covered_rows:
            continue
        if not is_meaningful_row(row_vals):
            continue
        if is_footer_row(a):
            continue

        parse_issues.append({
            "Source File": file_path.name,
            "Source Path": str(file_path),
            "Sheet Name": sheet_name,
            "Issue Type": "UNCOVERED_MEANINGFUL_ROW",
            "Message": "통합정리데이터의 Source Start/End 범위 어떤 곳에도 포함되지 않은 의미 있는 행입니다.",
            "Row No": chk_row,
            "Source Start Row": chk_row,
            "Source End Row": chk_row,
            "LINE NO": "",
            "ISO NO": "",
            "Item Name": row_vals[1] if len(row_vals) > 1 else "",
            "Row Snapshot": " | ".join(row_vals),
        })

    unique_lines = len(set([x["LINE NO"] for x in item_records if x["LINE NO"]]))
    unique_iso = len(set([x["ISO NO"] for x in item_records if x["ISO NO"]]))

    summary = {
        "Source File": file_path.name,
        "Source Path": str(file_path),
        "Sheet Name": sheet_name,
        "Rows": rows,
        "Columns": cols,
        "Item Count": len(item_records),
        "Unique LINE NO": unique_lines,
        "Unique ISO NO": unique_iso,
        "Excluded Header Count": len(excluded_headers),
        "Coverage Issue Count": len(parse_issues),
        "Status": "OK",
        "Message": "",
    }

    return item_records, excluded_headers, summary, parse_issues


# =========================
# EHT LIST 비교
# =========================
# EHT LIST 한 행(ISO 시트 + DN)을 TI 의 같은 ISO 시트(시트번호 붙은 ISO NO) + DN 품목과 비교한다.
# 길이는 파이프(Pipe/Труба/Tube)만, 밸브는 수량으로 두께별 합계를 낸다. (엘보/티/리듀서는 DN 이 없어 제외)
# TI Note 에 전기히팅 표시가 있는 품목이 있으면 그 품목만 비교한다. (EHT 길이는 히팅 구간 기준)
# 'TI 최소 두께' 는 비교한 TI 파이프·밸브 중 가장 얇은 두께 (EHT 에 보수적으로 반영할 값 = 열손실이 가장 큰 경우)
EHT_STATUSES = [  # (판정, 설명) - EHT비교요약 시트 순서
    ("일치", "TI 파이프·밸브 두께가 모두 EHT 두께와 같고, 길이 차이도 {tol:g}m 이내"),
    ("길이 차이", "두께는 같지만 EHT 길이와 TI 파이프 길이가 {tol:g}m 넘게 차이 남"),
    ("두께 혼재", "TI에 EHT 두께와 다른 두께가 섞여 있음 ('EHT 두께 비율' 확인)"),
    ("두께 불일치", "TI에 EHT 두께가 없음"),
    ("TI에 DN 없음", "TI에 ISO 시트는 있으나 해당 DN의 파이프/밸브가 없음"),
    ("TI에 ISO 없음", "TI-LST에 해당 ISO 시트가 없음"),
    ("TI 제외라인", "TI에서 제외(Excluded)된 라인에만 있음"),
    ("EHT에 없음", "TI Note에 전기히팅(There is electric heating) 표시가 있으나 EHT LIST에 해당 ISO 시트·DN이 없음"),
    ("비교 제외", "Pumping Trap, DN이 숫자 하나가 아님(40x20 등), ISO/두께 없음, 길이·밸브 0이고 TI에 DN 없음"),
]

EHT_TEMP_STATUSES = [  # (온도 판정, 설명) - 두께를 비교한 행(일치/길이 차이/두께 혼재/두께 불일치)만
    ("일치", "TI 온도가 EHT Operation Temp 와 같음 ('0.5 ~ 25' 는 범위 안, '37 / 290' 은 둘 중 하나와 같으면 일치)"),
    ("다름", "TI 온도가 EHT Operation Temp 와 다름"),
    ("비교 불가", "EHT 온도가 AMB·빈칸처럼 숫자가 아니거나 TI 온도가 없음 (온도 열을 못 찾으면 'EHT 열 위치' 확인)"),
]

EHT_REPORT_COLUMNS = [
    "판정", "내용",
    "ISO NO", "DN", "EHT 라인", "EHT 배관재질", "EHT 보온구분", "EHT Remark",
    "EHT 두께 (mm)", "TI 최소 두께 (mm)", "두께 변경 필요", "EHT 길이 (m)", "EHT 밸브 (ea)",
    "TI 파이프 두께별 길이", "TI 파이프 길이 (m)", "TI 중 EHT 두께 길이 (m)", "EHT 두께 비율 (%)", "길이 차이 (m)",
    "TI 밸브 두께별 수량", "TI 밸브 (ea)",
    "EHT 운전온도 (℃)", "TI 온도 (°C)", "온도 판정",
    "TI LINE NO", "TI 보온구분", "TI 파일",
    "EHT 파일", "EHT 행",
    "COMMENT",  # 비워 둠: 여기 적은 COMMENT 를 ti_comment_writer.py 가 TI-LST 에 옮겨 적음
]

EHT_SUMMARY_COLUMNS = ["판정", "건수", "설명"]

EHT_HEADER_SCAN_ROWS = 30  # 머리글('ISO NUMBER')을 찾을 시트 위쪽 행 수
EHT_FIELDS = [  # (항목, 안내용 이름, 머리글 글자) - 여러 줄 머리글을 합쳐 소문자로 비교, 위 항목·앞 글자부터 열을 정함
    ("ISO", "ISO NUMBER", ("iso number", "iso no")),
    ("DN", "Nominal Diameter", ("nominal diameter",)),
    ("THK", "Thickness", ("thickness", "thk")),
    ("Length", "Length", ("length",)),
    ("Valve", "No. of Valve", ("valve",)),
    ("Pipe Material", "pipe material", ("pipe material",)),
    ("Insulation", "Insulation material", ("insulation material", "material")),
    ("Operation Temp", "Operation Temp",  # ("oper", "temp") 처럼 묶은 글자는 둘 다 들어 있으면 맞음
     ("operation temp", "operating temp", "oper. temp", "운전온도", ("oper", "temp"), ("운전", "온도"))),
    ("Remark", "Piping Remark", ("remark",)),
    ("ID", "Equipment/piping number", ("piping number", "equipment")),
]
EHT_REQUIRED_FIELDS = ("ISO", "DN", "THK")
EHT_CHECK_FIELDS = {  # 열이 없거나 값이 비어 있으면 알려줄 항목 -> 못 하게 되는 비교
    "ISO": "비교 안 됨", "DN": "비교 안 됨", "THK": "비교 안 됨",
    "Length": "길이 비교 안 됨", "Valve": "밸브 수량 비교 안 됨",
    "Insulation": "INDOOR/OUTDOOR 구분 안 됨", "Operation Temp": "온도 비교 안 됨",
}
EHT_SUBHEADER_WORDS = ("nominal diameter", "thickness", "length", "material", "temp")
HEATING_NOTES = ("there is electric heating", "имеется электрообогрев")
INSULATION_TYPE_RE = re.compile(r"\b([A-Z]{1,3})\s*\(\s*([A-Z]{3,})\s*\)")  # H(OUTDOOR), H(REGEN) 등


def to_number(value):
    """숫자 또는 '1,15' / '0.1' 같은 글자를 float 로 바꾼다 (못 바꾸면 None)."""
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        num = float(value)
    else:
        try:
            num = float(clean_text(value).replace(" ", "").replace(",", "."))
        except ValueError:
            return None
    return num if math.isfinite(num) else None


def tidy_number(num):
    num = round(float(num), 3)  # 합계가 0 이면 int 로 들어옴 (Python 3.8 의 int 에는 is_integer 가 없음)
    return int(num) if num.is_integer() else num


def dn_number(value):
    """'DN 25' / 25 / '25' -> 25. 빈칸, '-', '40x20' 처럼 숫자 하나가 아니면 None."""
    if isinstance(value, str):
        value = re.sub(r"^\s*DN\s*", "", value, flags=re.IGNORECASE)
    num = to_number(value)
    return int(num) if num is not None and num.is_integer() else None


def normalize_iso(text):
    """비교용 ISO 키: 공백 제거, 대문자, 시트번호 2자리 ('...-ISO-00094-1' -> '...-ISO-00094-01')."""
    s = re.sub(r"\s+", "", clean_text(text)).upper()
    m = re.fullmatch(r"(.*-ISO-\d+)-(\d{1,2})", s)
    return f"{m.group(1)}-{int(m.group(2)):02d}" if m else s


def has_sheet_no(iso_key):
    return re.search(r"-ISO-\d+-\d{2}$", iso_key) is not None


def line_tokens(text):
    return {t for t in re.split(r"[-/\s()]+", clean_text(text).upper()) if t}


def value_order(value):
    """두께/온도 정렬: 숫자 -> 숫자가 아닌 글자(예: '30+30') -> 빈칸"""
    if isinstance(value, float):
        return (0, value, "")
    return (1, 0, value) if value else (2, 0, "")


def thk_label(thk):
    if thk is None:
        return "두께없음"
    return f"{thk:g}mm" if isinstance(thk, float) else f"{thk}mm"


def by_thk_text(amounts, unit):
    """{두께: 양} -> '20mm 2.00m, 30mm 1.00m' (두께 순)"""
    fmt = "{:.2f}m" if unit == "m" else "{:g}ea"
    return ", ".join(f"{thk_label(thk)} {fmt.format(amounts[thk])}" for thk in sorted(amounts, key=value_order))


def pipe_valve_text(pipe, valve):
    """파이프 길이 + 밸브 수량 -> '40mm 1.52m, 밸브 40mm 1ea'"""
    parts = [by_thk_text(pipe, "m")] if pipe else []
    if valve:
        parts.append("밸브 " + by_thk_text(valve, "ea"))
    return ", ".join(parts)


def distinct_temps(entries):
    """TI 품목 온도 목록 (숫자 먼저, 중복 없이)"""
    return sorted({e["temp"] for e in entries if e["temp"] is not None}, key=value_order)


def temp_ranges(value):
    """EHT 온도 -> [(최소, 최대)]. '70' -> [(70, 70)], '0.5 ~ 25' -> [(0.5, 25)],
    '37 / 290' -> [(37, 37), (290, 290)]. AMB 처럼 숫자가 아니면 None."""
    num = to_number(value)
    if num is not None:
        return [(num, num)]
    ranges = []
    for part in clean_text(value).replace("℃", "").replace("°C", "").split("/"):
        m = re.fullmatch(r"\s*(-?\d+(?:[.,]\d+)?)\s*~\s*(-?\d+(?:[.,]\d+)?)\s*", part)
        if m:
            ranges.append(tuple(sorted((to_number(m.group(1)), to_number(m.group(2))))))
            continue
        num = to_number(part)
        if num is None:
            return None
        ranges.append((num, num))
    return ranges


def compare_temp(eht_temp, ti_temps):
    """EHT Operation Temp 와 TI 온도 비교 -> '일치' / '다름' / '비교 불가'"""
    ranges = temp_ranges(eht_temp)
    numbers = [t for t in ti_temps if isinstance(t, float)]
    if not ranges or not numbers or len(numbers) < len(ti_temps):
        return "비교 불가"
    return "일치" if all(any(lo <= t <= hi for lo, hi in ranges) for t in numbers) else "다름"


def header_text(value):
    """머리글 비교용 글자: 전각 문자·특수 공백·보이지 않는 문자를 정리하고 소문자로 ('ISO  NUMBER' -> 'iso number')"""
    text = unicodedata.normalize("NFKC", clean_text(value))
    return " ".join("".join(ch for ch in text if unicodedata.category(ch) != "Cf").lower().split())


def is_header_row(row, iso_col):
    """머리글 바로 아래 줄도 머리글인지: 머리글 글자가 있거나, 숫자 칸이 없고 ISO 칸이 비어 있으면(병합) 머리글"""
    texts = [header_text(v) for v in row]
    if any(word in t for t in texts for word in EHT_SUBHEADER_WORDS):
        return True
    has_number = any(isinstance(v, (int, float)) and not isinstance(v, bool) for v in row)
    return not has_number and not (texts[iso_col] if iso_col < len(texts) else "")


def find_eht_columns(head):
    """EHT LIST 머리글을 찾아 ({항목: [후보 열 번호]}, 머리글 첫 줄, 데이터 시작 위치, 열별 머리글)을 돌려준다.
    없으면 None."""
    for idx, row in enumerate(head):
        labels = [header_text(v) for v in row]
        iso_cols = [col for col, t in enumerate(labels) if "iso number" in t or "iso no" in t]
        if not iso_cols:
            continue

        # 여러 줄 머리글(병합)이면 아래 줄 글자도 합침: 'Insulation' + 'material', 'Process data' + 'Operation Temp (℃)'
        start = idx + 1
        while start < len(head) and start - idx <= 3 and is_header_row(head[start], iso_cols[0]):
            sub = [header_text(v) for v in head[start]]
            labels = [f"{a} {b}".strip() for a, b in zip_longest(labels, sub, fillvalue="")]
            start += 1

        candidates = {}
        for field, _, words in EHT_FIELDS:
            found = [col for word in words for col, label in enumerate(labels)
                     if (all(w in label for w in word) if isinstance(word, tuple) else word in label)]
            if found:
                candidates[field] = list(dict.fromkeys(found))
        if all(field in candidates for field in EHT_REQUIRED_FIELDS):
            return candidates, idx, start, labels
    return None


def formula_columns(file_path, sheet_name, data_start, columns):
    """값이 비어 있는 열 중 수식이 들어 있는 열 (수식의 계산 값이 파일에 저장되지 않은 경우)"""
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=False)
    try:
        ws = wb[sheet_name]
        ws.reset_dimensions()
        found = set()
        for row in islice(ws.iter_rows(values_only=True), data_start, None):
            found.update(col for col in columns
                         if col < len(row) and isinstance(row[col], str) and row[col].startswith("="))
            if found == set(columns):
                break
        return found
    finally:
        wb.close()


def read_eht_list(file_path):
    """EHT LIST 를 읽어 (데이터 행, 읽은 정보)를 돌려준다.
    'ISO NUMBER' 머리글이 있는 시트가 없으면 None (= EHT LIST 아님).

    머리글이 맞는 시트가 여럿이면 찾은 항목이 가장 많은 시트를 읽는다.
    같은 머리글 열이 여럿이면 앞의 열을 쓰되, 값이 하나도 없는 열은 건너뛴다.
    """
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)  # VLOOKUP 등 수식은 계산된 값으로 읽음
    try:
        best = None
        for ws in wb.worksheets:
            ws.reset_dimensions()
            head = [list(row) for row in islice(ws.iter_rows(values_only=True), EHT_HEADER_SCAN_ROWS)]
            found = find_eht_columns(head)
            if found and (best is None or len(found[0]) > len(best[1][0])):
                best = (ws, found)
        if best is None:
            return None
        ws, (candidates, header_row, data_start, labels) = best
        sheet_name = ws.title
        data = [list(row) for row in islice(ws.iter_rows(values_only=True), data_start, None)]
    finally:
        wb.close()

    def has_values(col):
        return any(col < len(row) and row[col] not in (None, "") for row in data)

    cols = {}
    for field, _, _ in EHT_FIELDS:
        options = [col for col in candidates.get(field, []) if col not in cols.values()]
        if options:
            cols[field] = next((col for col in options if has_values(col)), options[0])
    if "ID" not in cols and 0 not in cols.values():
        cols["ID"] = 0  # 첫 열 = Equipment/piping number

    records = []
    for row_no, row in enumerate(data, start=data_start + 1):
        rec = {"EHT 파일": file_path.name, "EHT 행": row_no}
        for field, col in cols.items():
            value = row[col] if col < len(row) else None
            rec[field] = clean_text(value) if isinstance(value, str) else value
        if any(rec.get(field) not in (None, "") for field in ("ID", "DN", "ISO")):
            records.append(rec)

    empty = [field for field in EHT_CHECK_FIELDS if field in cols and not has_values(cols[field])]
    formulas = formula_columns(file_path, sheet_name, data_start, [cols[f] for f in empty]) if empty else set()
    info = {
        "file": file_path.name,
        "sheet": sheet_name,
        "header_rows": f"{header_row + 1}~{data_start}행" if data_start - header_row > 1 else f"{data_start}행",
        "columns": {field: get_column_letter(col + 1) for field, col in cols.items()},
        "missing": [field for field in EHT_CHECK_FIELDS if field not in cols],
        "empty": empty,
        "formula": [field for field in empty if cols[field] in formulas],
        "labels": labels,
    }
    return records, info


def eht_column_text(info):
    """'EHT LIST.xlsx' 'EHT LIST' 시트 1~2행: ISO NUMBER=Q, Nominal Diameter=C, ..."""
    names = {field: name for field, name, _ in EHT_FIELDS}
    columns = ", ".join(f"{names[field]}={letter}" for field, letter in info["columns"].items())
    return f"{info['file']} '{info['sheet']}' 시트, 머리글 {info['header_rows']}: {columns}"


def eht_column_problems(info):
    """EHT LIST 에서 찾지 못했거나 값이 비어 있는 열 안내 문구"""
    names = {field: name for field, name, _ in EHT_FIELDS}
    problems = [f"'{names[field]}' 열을 찾지 못함 → {EHT_CHECK_FIELDS[field]}" for field in info["missing"]]
    for field in info["empty"]:
        where = f"'{names[field]}' 열({info['columns'][field]}열) 값이 모두 비어 있음"
        if field in info["formula"]:
            where += " - 수식의 계산 값이 파일에 저장되어 있지 않음 (엑셀에서 열어 저장한 뒤 다시 실행)"
        problems.append(f"{where} → {EHT_CHECK_FIELDS[field]}")
    return problems


def build_ti_groups(all_items):
    """TI 품목을 (ISO 시트, DN) 별로 묶는다. DN 이 있는 품목(파이프/튜브/밸브)만 들어감."""
    groups = defaultdict(list)
    iso_dns = {}  # ISO 시트 -> TI 에 있는 DN (DN 없는 품목만 있어도 ISO 시트는 등록)
    for rec in all_items:
        iso = normalize_iso(rec.get("ISO NO"))
        if not iso:
            continue
        dns = iso_dns.setdefault(iso, set())
        dn = dn_number(rec.get("Nominal Diameter"))
        if dn is None:
            continue
        dns.add(dn)

        is_valve = "valve" in rec["Item Name"].lower()
        thk_text = rec["Insulation Thickness (mm)"] or rec["Insulation Detail Thickness (mm)"]
        thk = to_number(thk_text)
        temp = to_number(rec["Temperature (°C)"])
        qty = to_number(rec["Qty"])
        m_type = INSULATION_TYPE_RE.search(rec["Line Header"])
        groups[(iso, dn)].append({
            "iso_text": rec["ISO NO"],
            "row": rec["Source Start Row"],
            "line": rec["LINE NO"],
            "type": f"{m_type.group(1)}({m_type.group(2)})" if m_type else "",
            "file": rec["Source File"],
            "valve": is_valve,
            "thk": round(thk, 1) if thk is not None else (thk_text or None),  # 숫자가 아니면 글자 그대로
            "temp": round(temp, 1) if temp is not None else (rec["Temperature (°C)"] or None),
            "length": 0.0 if is_valve else (to_number(rec["Length/Height (m)"]) or 0.0),
            "qty": (1.0 if qty is None else qty) if is_valve else 0.0,
            "excluded": rec["Excluded"] == "Y",
            "heating": any(word in rec["Note"].lower() for word in HEATING_NOTES),
        })
    return groups, iso_dns


def summarize_ti(entries):
    """TI 품목의 두께별 파이프 길이/밸브 수량과 비교 시트에 넣을 칸을 돌려준다."""
    pipe, valve = defaultdict(float), defaultdict(float)
    for e in entries:
        if e["valve"]:
            valve[e["thk"]] += e["qty"]
        else:
            pipe[e["thk"]] += e["length"]
    numeric_thks = [t for t in list(pipe) + list(valve) if isinstance(t, float)]
    columns = {
        "TI 최소 두께 (mm)": tidy_number(min(numeric_thks)) if numeric_thks else None,
        "TI 파이프 두께별 길이": by_thk_text(pipe, "m"),
        "TI 파이프 길이 (m)": tidy_number(sum(pipe.values())),
        "TI 밸브 두께별 수량": by_thk_text(valve, "ea"),
        "TI 밸브 (ea)": tidy_number(sum(valve.values())),
        "TI 온도 (°C)": ", ".join(f"{t:g}" if isinstance(t, float) else t for t in distinct_temps(entries)),
        "TI LINE NO": ", ".join(dict.fromkeys(e["line"] for e in entries)),
        "TI 보온구분": ", ".join(dict.fromkeys(e["type"] for e in entries if e["type"])),
        "TI 파일": ", ".join(dict.fromkeys(e["file"] for e in entries)),
    }
    return pipe, valve, columns


def compare_eht_row(row, iso, dn, groups, ti_iso_dns, split, out):
    """EHT 한 행을 TI 와 비교해 (판정, 내용 목록)을 돌려주고 TI 칸은 out 에 채운다."""
    thk = to_number(row.get("THK"))
    length = to_number(row.get("Length")) or 0.0
    valves = to_number(row.get("Valve")) or 0.0

    if "pumping trap" in clean_text(row.get("Remark")).lower():
        return "비교 제외", ["Pumping Trap (장비)"]
    if dn is None:
        return "비교 제외", [f"DN이 숫자 하나가 아님 ({clean_text(row.get('DN')) or '빈칸'})"]
    if not iso:
        return "비교 제외", ["ISO NUMBER 없음"]
    if thk is None:
        return "비교 제외", [f"EHT 두께 확인 ({clean_text(row.get('THK')) or '빈칸'})"]
    if iso not in ti_iso_dns:
        if re.sub(r"-\d{2}$", "", iso) in ti_iso_dns:
            return "TI에 ISO 없음", ["TI에는 시트번호 없는 ISO NO로 있음 (에러로그 SHEET_NO_NOT_FOUND 확인)"]
        return "TI에 ISO 없음", []

    entries = groups.get((iso, dn), [])
    if not entries:
        if length == 0 and valves == 0:
            return "비교 제외", [f"EHT 길이·밸브 0 (플랜지/서포트만), TI에 DN {dn} 없음"]
        ti_dns = ", ".join(map(str, sorted(ti_iso_dns[iso]))) or "없음"
        return "TI에 DN 없음", [f"이 ISO 시트의 TI DN: {ti_dns}"]

    active = [e for e in entries if not e["excluded"]]
    if not active:
        out.update(summarize_ti(entries)[2])
        return "TI 제외라인", []

    notes = []
    # 같은 ISO 시트에 라인이 여럿이면 EHT 라인번호(끝의 시트번호 제외)가 들어 있는 TI 라인만 비교
    eht_line = line_tokens(re.sub(r"-\d{1,2}$", "", clean_text(row.get("ID"))))
    if eht_line:
        same_line = [e for e in active if eht_line <= line_tokens(e["line"])]
        if not same_line:
            notes.append("TI LINE NO에 EHT 라인번호가 없음")
        elif len(same_line) < len(active):
            active = same_line
    # EHT 가 같은 ISO 시트·DN 을 보온구분(H(INDOOR)/H(OUTDOOR) 등)으로 나눠 적었으면 TI 도 같은 구분끼리 비교
    if split:
        eht_type = re.sub(r"\s+", "", clean_text(row.get("Insulation"))).upper()
        same_type = [e for e in active if e["type"] == eht_type]
        if same_type:
            active = same_type
        else:
            notes.append("TI를 보온구분으로 나눌 수 없어 해당 DN 전체와 비교")
    # EHT 는 히팅 구간 기준이므로, TI Note 에 전기히팅 표시가 있는 품목이 있으면 그 품목끼리만 비교
    heated = [e for e in active if e["heating"]]
    if heated and len(heated) < len(active):
        pipe_x, valve_x, _ = summarize_ti([e for e in active if not e["heating"]])
        notes.append("전기히팅 표시 없는 TI 품목 제외: " + pipe_valve_text(pipe_x, valve_x))
        active = heated

    pipe, valve, columns = summarize_ti(active)
    out.update(columns)
    thk = round(thk, 1)
    ti_length = sum(pipe.values())
    ti_valves = round(sum(valve.values()), 3)
    diff = round(ti_length - length, 3)
    out["TI 중 EHT 두께 길이 (m)"] = tidy_number(pipe.get(thk, 0.0))
    if ti_length > 0:
        out["EHT 두께 비율 (%)"] = round(pipe.get(thk, 0.0) / ti_length * 100, 1)
    elif ti_valves > 0:
        out["EHT 두께 비율 (%)"] = round(valve.get(thk, 0.0) / ti_valves * 100, 1)  # 파이프가 없으면 밸브 수량 기준
    out["길이 차이 (m)"] = tidy_number(diff)

    thks = set(pipe) | set(valve)
    if thk not in thks:
        status = "두께 불일치"
        notes.insert(0, "TI 두께: " + ", ".join(thk_label(t) for t in sorted(thks, key=value_order)))
    elif len(thks) > 1:
        status = "두께 혼재"
        other_pipe = {t: v for t, v in pipe.items() if t != thk}
        other_valve = {t: v for t, v in valve.items() if t != thk}
        notes.insert(0, "EHT 두께가 아닌 부분: " + pipe_valve_text(other_pipe, other_valve))
    elif abs(diff) > EHT_LENGTH_TOLERANCE:
        status = "길이 차이"
    else:
        status = "일치"

    if abs(diff) > EHT_LENGTH_TOLERANCE:
        notes.append(f"길이 TI {ti_length:.2f}m vs EHT {length:g}m")
    if ti_valves != round(valves, 3):
        notes.append(f"밸브 TI {ti_valves:g}ea vs EHT {valves:g}ea")

    # EHT 두께가 TI 최소 두께(가장 얇은 = 열손실이 가장 큰 경우)와 다르면 TI 최소 두께로 반영해야 함
    ti_min = columns["TI 최소 두께 (mm)"]
    if ti_min is not None:
        out["두께 변경 필요"] = "Y" if float(ti_min) != thk else "N"
    out["온도 판정"] = compare_temp(row.get("Operation Temp"), distinct_temps(active))
    return status, notes


def compare_eht(eht_rows, all_items, eht_infos=()):
    """EHT LIST 각 행을 같은 ISO 시트·DN 의 TI 두께/길이와 비교한다. (비교 행, 요약 행)을 돌려준다.
    eht_infos: read_eht_list 가 돌려준 읽은 정보 (요약 시트에 EHT 열 위치를 적음)"""
    groups, ti_iso_dns = build_ti_groups(all_items)

    prepared = []
    eht_keys = Counter()
    for row in eht_rows:
        iso = normalize_iso(row.get("ISO"))
        if iso and not has_sheet_no(iso):  # EHT ISO 에 시트번호가 없으면 라인번호 끝 숫자로 보충
            sheet_no = get_sheet_no(row.get("ID"))
            if sheet_no:
                iso = normalize_iso(f"{iso}-{sheet_no}")
        dn = dn_number(row.get("DN"))
        prepared.append((row, iso, dn))
        if iso and dn is not None:
            eht_keys[(iso, dn)] += 1
    eht_isos = {iso for _, iso, _ in prepared if iso}

    report = []
    for row, iso, dn in prepared:
        out = dict.fromkeys(EHT_REPORT_COLUMNS)
        out.update({
            "ISO NO": row.get("ISO"), "DN": row.get("DN"), "EHT 라인": row.get("ID"),
            "EHT 배관재질": row.get("Pipe Material"), "EHT 보온구분": row.get("Insulation"),
            "EHT Remark": row.get("Remark"), "EHT 두께 (mm)": row.get("THK"),
            "EHT 길이 (m)": row.get("Length"), "EHT 밸브 (ea)": row.get("Valve"),
            "EHT 운전온도 (℃)": row.get("Operation Temp"),
            "EHT 파일": row["EHT 파일"], "EHT 행": row["EHT 행"],
        })
        status, notes = compare_eht_row(row, iso, dn, groups, ti_iso_dns, eht_keys[(iso, dn)] > 1, out)
        out["판정"] = status
        out["내용"] = " / ".join(notes)
        report.append(out)

    # TI 에는 전기히팅 표시가 있는데 EHT LIST 에 없는 ISO 시트·DN (히팅 표시 품목만 보여줌)
    for (iso, dn), entries in sorted(groups.items()):
        heated = [e for e in entries if e["heating"] and not e["excluded"]]
        if (iso, dn) in eht_keys or not heated:
            continue
        out = dict.fromkeys(EHT_REPORT_COLUMNS)
        out.update(summarize_ti(heated)[2])
        note = f"EHT LIST에 이 ISO 시트는 있으나 DN {dn} 없음" if iso in eht_isos else "EHT LIST에 이 ISO 시트가 없음"
        if not has_sheet_no(iso):
            note += " (TI ISO NO에 시트번호 없음)"
        out.update({"판정": "EHT에 없음", "내용": note, "ISO NO": heated[0]["iso_text"], "DN": dn})
        report.append(out)

    counts = Counter(rec["판정"] for rec in report)
    summary = [{"판정": status, "건수": counts[status], "설명": desc.format(tol=EHT_LENGTH_TOLERANCE)}
               for status, desc in EHT_STATUSES]
    summary.append({"판정": "합계", "건수": len(report),
                    "설명": f"EHT LIST {len(eht_rows)}행 + EHT에 없음 {counts['EHT에 없음']}건"})
    summary.append({"판정": "두께 변경 필요", "건수": sum(rec["두께 변경 필요"] == "Y" for rec in report),
                    "설명": "EHT 두께가 TI 최소 두께와 다름 → TI 최소 두께(가장 얇은 = 열손실 최대)로 반영"})
    temp_counts = Counter(rec["온도 판정"] for rec in report if rec["온도 판정"])
    summary.extend({"판정": f"온도 {status}", "건수": temp_counts[status], "설명": desc}
                   for status, desc in EHT_TEMP_STATUSES)
    summary.append({"판정": "비교 기준", "설명": (
        "ISO 시트(ISO NO + 시트번호)와 DN이 같은 TI 파이프(길이)·밸브(수량)의 보온 두께와 비교. "
        "TI Note에 전기히팅 표시가 있는 품목이 있으면 그 품목만 비교. 엘보/티/리듀서는 DN이 없어 제외. "
        "TI 최소 두께 = 비교한 TI 파이프·밸브 중 가장 얇은 두께. "
        "온도는 EHT Operation Temp 와 TI Temperature 비교.")})
    for info in eht_infos:
        problems = eht_column_problems(info)
        summary.append({"판정": "EHT 열 위치", "설명": eht_column_text(info) + "".join(f" / !! {p}" for p in problems)})
    return report, summary


# =========================
# 파일 처리 (병렬)
# =========================
def process_file(file_path):
    """파일 1개 처리 (병렬 처리 단위). 오류가 나면 예외 대신 ERROR 요약/로그를 돌려준다."""
    try:
        return parse_insulation_workbook(file_path)
    except Exception as e:
        summary = {
            "Source File": file_path.name,
            "Source Path": str(file_path),
            "Sheet Name": "",
            "Rows": "",
            "Columns": "",
            "Item Count": 0,
            "Unique LINE NO": 0,
            "Unique ISO NO": 0,
            "Excluded Header Count": 0,
            "Coverage Issue Count": 0,
            "Status": "ERROR",
            "Message": str(e),
        }
        error = {
            "Source File": file_path.name,
            "Source Path": str(file_path),
            "Sheet Name": "",
            "Issue Type": "FILE_PROCESSING_ERROR",
            "Message": "파일 처리 중 예외 발생",
            "Row No": "",
            "Source Start Row": "",
            "Source End Row": "",
            "LINE NO": "",
            "ISO NO": "",
            "Item Name": "",
            "Row Snapshot": "",
            "Error": str(e),
            "Traceback": traceback.format_exc(),
        }
        return [], [], summary, [error]


def print_progress(idx, total, file_path, result):
    items, excluded, summary, parse_issues = result
    if summary["Status"] == "ERROR":
        print(f"[{idx}/{total}] !! 오류 발생: {file_path.name}")
        print(f"       {summary['Message']}")
    else:
        print(f"[{idx}/{total}] {file_path.name} -> "
              f"ITEM {len(items)}건, 제외라인 {len(excluded)}건, 검토이슈 {len(parse_issues)}건")


def process_files(target_files):
    """모든 파일을 처리해 파일 순서대로 결과 목록을 돌려준다 (가능하면 여러 프로세스로 동시 처리)."""
    total = len(target_files)
    workers = MAX_WORKERS if MAX_WORKERS > 0 else min(os.cpu_count() or 1, 8)
    workers = min(workers, total)

    if workers > 1:
        print(f"병렬 처리: 프로세스 {workers}개")
        try:
            results = []
            with ProcessPoolExecutor(max_workers=workers) as executor:
                for idx, (file_path, result) in enumerate(
                    zip(target_files, executor.map(process_file, target_files)), start=1
                ):
                    print_progress(idx, total, file_path, result)
                    results.append(result)
            return results
        except Exception as e:
            # 일부 IDE(Spyder/Jupyter 등) 실행 환경에서는 병렬 처리가 막힐 수 있음
            print(f"병렬 처리를 사용할 수 없어 순차 처리로 다시 진행합니다. ({e!r})")

    results = []
    for idx, file_path in enumerate(target_files, start=1):
        result = process_file(file_path)
        print_progress(idx, total, file_path, result)
        results.append(result)
    return results


# =========================
# 결과 엑셀 저장
# =========================
ITEM_COLUMNS = [
    "Source File", "Section",
    "LINE NO", "ISO NO", "INDOOR/OUTDOOR", "Excluded", "Exclusion Letter",
    "Item Name", "Brand/Position",
    "Size (mm)", "Nominal Diameter", "Length/Height (m)", "Qty", "Temperature (°C)",
    "Purpose/Location",
    "Insulation Material", "Insulation Thickness (mm)",
    "Insulation Detail", "Insulation Detail Thickness (mm)",
    "Covering Material", "Covering Thickness (mm)", "Covering Surface (m²)",
    "Surface (m²)", "Volume of Insulation Layer (m³)",
    "Document Designation", "Note",
    "Line Header",
    "Sheet Name", "Source Path",
    "Source Start Row", "Source End Row",
]

EXCLUDED_COLUMNS = [
    "Source File", "Section", "LINE NO", "ISO NO", "INDOOR/OUTDOOR",
    "Excluded", "Exclusion Letter", "Line Header", "Sheet Name", "Source Path",
    "Source Start Row", "Source End Row",
]

ERROR_COLUMNS = [
    "Source File", "Source Path", "Sheet Name", "Issue Type", "Message",
    "Row No", "Source Start Row", "Source End Row",
    "LINE NO", "ISO NO", "Item Name", "Row Snapshot",
    "Error", "Traceback",
]


def order_columns(records, preferred_cols=()):
    """레코드에 있는 컬럼을 preferred_cols 순서로 앞에 두고, 나머지는 처음 나온 순서대로 뒤에 붙인다."""
    all_cols = list(dict.fromkeys(key for rec in records for key in rec))
    existing_cols = [c for c in preferred_cols if c in all_cols]
    other_cols = [c for c in all_cols if c not in existing_cols]
    return existing_cols + other_cols


def column_widths(columns, rows, min_width=10, max_width=45):
    """열 너비 자동 조정: 헤더 포함 최대 글자 수 × 0.95 를 min~max 범위로 제한."""
    max_lens = [len(str(col)) for col in columns]
    for row in rows:
        for idx, value in enumerate(row):
            if value is None:
                continue
            length = len(value) if isinstance(value, str) else len(str(value))
            if length > max_lens[idx]:
                max_lens[idx] = length
    return [min(max(length * 0.95, min_width), max_width) for length in max_lens]


# ----- xlsx 직접 작성 -----
# openpyxl 로 셀마다 서식을 지정해 저장하면 수십만 셀에서 매우 느려서(원래 코드의 최대 병목),
# 필요한 서식 2가지만 정의한 xlsx(XML 묶음)를 직접 만든다.
#   헤더: 굵은 흰 글씨 + 남색(1F4E78) 배경 + 가운데 정렬 + 줄바꿈
#   데이터: 위쪽 정렬 + 줄바꿈 / 틀 고정(A2), 자동 필터, 열 너비 자동
HEADER_STYLE = 1
DATA_STYLE = 2

EXCEL_MAX_CELL_CHARS = 32767
ILLEGAL_XML_CHARS_RE = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")

XML_DECLARATION = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
NS_MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
NS_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_PACKAGE_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml"

STYLES_XML = (
    XML_DECLARATION
    + f'<styleSheet xmlns="{NS_MAIN}">'
    '<fonts count="2">'
    '<font><sz val="11"/><name val="Calibri"/><family val="2"/></font>'
    '<font><b/><sz val="11"/><color rgb="FFFFFFFF"/><name val="Calibri"/><family val="2"/></font>'
    '</fonts>'
    '<fills count="3">'
    '<fill><patternFill patternType="none"/></fill>'
    '<fill><patternFill patternType="gray125"/></fill>'
    '<fill><patternFill patternType="solid"><fgColor rgb="FF1F4E78"/><bgColor indexed="64"/></patternFill></fill>'
    '</fills>'
    '<borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>'
    '<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
    '<cellXfs count="3">'
    '<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>'
    '<xf numFmtId="0" fontId="1" fillId="2" borderId="0" xfId="0" applyFont="1" applyFill="1" applyAlignment="1">'
    '<alignment horizontal="center" vertical="center" wrapText="1"/></xf>'
    '<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0" applyAlignment="1">'
    '<alignment vertical="top" wrapText="1"/></xf>'
    '</cellXfs>'
    '<cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles>'
    '</styleSheet>'
)


def xml_escape(text):
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;")


class SharedStrings:
    """xlsx 공유 문자열 표: 같은 문자열은 한 번만 저장하고 셀에는 번호만 기록한다."""

    def __init__(self):
        self.index_of = {}
        self.ref_count = 0

    def index(self, text):
        self.ref_count += 1
        idx = self.index_of.get(text)
        if idx is None:
            idx = self.index_of[text] = len(self.index_of)
        return idx

    def to_xml(self):
        items = "".join(
            '<si><t xml:space="preserve">'
            + xml_escape(ILLEGAL_XML_CHARS_RE.sub("", text)[:EXCEL_MAX_CELL_CHARS])
            + "</t></si>"
            for text in self.index_of
        )
        return (XML_DECLARATION
                + f'<sst xmlns="{NS_MAIN}" count="{self.ref_count}" uniqueCount="{len(self.index_of)}">'
                + items + "</sst>")


def cell_xml(ref, value, style, strings):
    style_attr = f' s="{style}"' if style else ""
    if value is None or value == "":
        return f'<c r="{ref}"{style_attr}/>' if style else ""
    if isinstance(value, bool):
        return f'<c r="{ref}"{style_attr} t="b"><v>{int(value)}</v></c>'
    if isinstance(value, (int, float)):
        if not math.isfinite(value):
            return f'<c r="{ref}"{style_attr}/>'
        return f'<c r="{ref}"{style_attr}><v>{value!r}</v></c>'
    return f'<c r="{ref}"{style_attr} t="s"><v>{strings.index(str(value))}</v></c>'


def write_sheet_xml(fh, columns, rows, strings, selected):
    tab = ' tabSelected="1"' if selected else ""
    head = XML_DECLARATION + f'<worksheet xmlns="{NS_MAIN}" xmlns:r="{NS_REL}">'

    if not rows:
        fh.write((
            head
            + '<dimension ref="A1"/>'
            + f'<sheetViews><sheetView{tab} workbookViewId="0"/></sheetViews>'
            + '<sheetFormatPr defaultRowHeight="15"/>'
            + '<sheetData><row r="1">' + cell_xml("A1", "데이터 없음", 0, strings) + "</row></sheetData>"
            + "</worksheet>"
        ).encode("utf-8"))
        return

    letters = [get_column_letter(idx) for idx in range(1, len(columns) + 1)]
    data_range = f"A1:{letters[-1]}{len(rows) + 1}"
    cols_xml = "".join(
        f'<col min="{idx}" max="{idx}" width="{width:.2f}" customWidth="1"/>'
        for idx, width in enumerate(column_widths(columns, rows), start=1)
    )
    header_cells = "".join(
        cell_xml(f"{letter}1", col, HEADER_STYLE, strings) for letter, col in zip(letters, columns)
    )
    fh.write((
        head
        + f'<dimension ref="{data_range}"/>'
        + f'<sheetViews><sheetView{tab} workbookViewId="0">'
        + '<pane ySplit="1" topLeftCell="A2" activePane="bottomLeft" state="frozen"/>'
        + '<selection pane="bottomLeft" activeCell="A2" sqref="A2"/>'
        + "</sheetView></sheetViews>"
        + '<sheetFormatPr defaultRowHeight="15"/>'
        + f"<cols>{cols_xml}</cols>"
        + f'<sheetData><row r="1">{header_cells}</row>'
    ).encode("utf-8"))

    chunk = []
    for r_idx, row in enumerate(rows, start=2):
        cells = "".join(
            cell_xml(f"{letter}{r_idx}", value, DATA_STYLE, strings) for letter, value in zip(letters, row)
        )
        chunk.append(f'<row r="{r_idx}">{cells}</row>')
        if len(chunk) >= 2000:
            fh.write("".join(chunk).encode("utf-8"))
            chunk = []
    chunk.append(f'</sheetData><autoFilter ref="{data_range}"/></worksheet>')
    fh.write("".join(chunk).encode("utf-8"))


def write_xlsx(output_path, sheets):
    """sheets: [(시트명, 컬럼 목록, 행 목록)] 을 xlsx 파일 하나로 저장한다."""
    strings = SharedStrings()
    count = len(sheets)

    sheet_entries = []
    filter_names = []
    for idx, (title, columns, rows) in enumerate(sheets, start=1):
        sheet_entries.append(f'<sheet name="{xml_escape(title)}" sheetId="{idx}" r:id="rId{idx}"/>')
        if rows:
            last_ref = f"${get_column_letter(len(columns))}${len(rows) + 1}"
            quoted_title = "'" + title.replace("'", "''") + "'"
            filter_names.append(
                f'<definedName name="_xlnm._FilterDatabase" localSheetId="{idx - 1}" hidden="1">'
                f"{xml_escape(quoted_title)}!$A$1:{last_ref}</definedName>"
            )

    workbook_xml = (
        XML_DECLARATION
        + f'<workbook xmlns="{NS_MAIN}" xmlns:r="{NS_REL}">'
        + '<bookViews><workbookView activeTab="0"/></bookViews>'
        + "<sheets>" + "".join(sheet_entries) + "</sheets>"
        + ("<definedNames>" + "".join(filter_names) + "</definedNames>" if filter_names else "")
        + "</workbook>"
    )
    workbook_rels = (
        XML_DECLARATION
        + f'<Relationships xmlns="{NS_PACKAGE_REL}">'
        + "".join(
            f'<Relationship Id="rId{idx}" Type="{NS_REL}/worksheet" Target="worksheets/sheet{idx}.xml"/>'
            for idx in range(1, count + 1)
        )
        + f'<Relationship Id="rId{count + 1}" Type="{NS_REL}/styles" Target="styles.xml"/>'
        + f'<Relationship Id="rId{count + 2}" Type="{NS_REL}/sharedStrings" Target="sharedStrings.xml"/>'
        + "</Relationships>"
    )
    content_types = (
        XML_DECLARATION
        + '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        + '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        + '<Default Extension="xml" ContentType="application/xml"/>'
        + f'<Override PartName="/xl/workbook.xml" ContentType="{CONTENT_TYPE}.sheet.main+xml"/>'
        + f'<Override PartName="/xl/styles.xml" ContentType="{CONTENT_TYPE}.styles+xml"/>'
        + f'<Override PartName="/xl/sharedStrings.xml" ContentType="{CONTENT_TYPE}.sharedStrings+xml"/>'
        + "".join(
            f'<Override PartName="/xl/worksheets/sheet{idx}.xml" ContentType="{CONTENT_TYPE}.worksheet+xml"/>'
            for idx in range(1, count + 1)
        )
        + "</Types>"
    )
    root_rels = (
        XML_DECLARATION
        + f'<Relationships xmlns="{NS_PACKAGE_REL}">'
        + f'<Relationship Id="rId1" Type="{NS_REL}/officeDocument" Target="xl/workbook.xml"/>'
        + "</Relationships>"
    )

    with zipfile.ZipFile(output_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", content_types)
        zf.writestr("_rels/.rels", root_rels)
        zf.writestr("xl/workbook.xml", workbook_xml)
        zf.writestr("xl/_rels/workbook.xml.rels", workbook_rels)
        zf.writestr("xl/styles.xml", STYLES_XML)
        for idx, (title, columns, rows) in enumerate(sheets, start=1):
            with zf.open(f"xl/worksheets/sheet{idx}.xml", "w") as fh:
                write_sheet_xml(fh, columns, rows, strings, selected=(idx == 1))
        zf.writestr("xl/sharedStrings.xml", strings.to_xml())


def save_output_excel(output_path, all_items, all_excluded, summaries, errors, eht_sheets=()):
    """eht_sheets: EHT LIST 가 있을 때 뒤에 붙일 [(시트명, 레코드, 컬럼 순서)]"""
    sheets = []
    for title, records, preferred_cols in [
        ("통합정리데이터", all_items, ITEM_COLUMNS),
        ("파일별요약", summaries, ()),
        ("제외라인목록", all_excluded, EXCLUDED_COLUMNS),
        ("에러로그", errors, ERROR_COLUMNS),
        *eht_sheets,
    ]:
        columns = order_columns(records, preferred_cols)
        rows = [[rec.get(col) for col in columns] for rec in records]
        sheets.append((title, columns, rows))

    write_xlsx(output_path, sheets)


def main():
    print("=== START PROGRAM ===")
    started = time.perf_counter()

    base_dir = Path(__file__).resolve().parent
    print(f"작업 폴더: {base_dir}")

    files = list(base_dir.rglob("*.xlsx")) if RECURSIVE else list(base_dir.glob("*.xlsx"))

    target_files = []
    for f in files:
        name = f.name
        if any(keyword.lower() in name.lower() for keyword in SKIP_KEYWORDS):
            continue
        target_files.append(f)

    all_items = []
    all_excluded = []
    summaries = []
    errors = []

    # EHT LIST: 파일명에 EHT_FILE_KEYWORD 가 있고 'ISO NUMBER' 머리글이 있으면 TI-LST 대신 EHT LIST 로 읽음
    eht_files = []
    eht_rows = []
    eht_infos = []
    for f in sorted(target_files):
        if not EHT_FILE_KEYWORD or EHT_FILE_KEYWORD.lower() not in f.name.lower():
            continue
        try:
            result = read_eht_list(f)
        except Exception as e:
            errors.append({
                "Source File": f.name, "Source Path": str(f), "Sheet Name": "",
                "Issue Type": "EHT_LIST_ERROR", "Message": "EHT LIST 파일을 읽는 중 예외 발생",
                "Error": str(e), "Traceback": traceback.format_exc(),
            })
            print(f"!! EHT LIST 읽기 오류: {f.name} ({e})")
            target_files.remove(f)
            continue
        if result is None:
            continue
        rows, info = result
        eht_files.append(f)
        eht_rows.extend(rows)
        eht_infos.append(info)
        print(f"EHT LIST: {f.name} ({len(rows)}행)")
        print(f"  읽은 열: {eht_column_text(info)}")
        for problem in eht_column_problems(info):  # 못 찾은 열/빈 열은 화면과 에러로그에 남김
            print(f"  !! {problem}")
            errors.append({
                "Source File": f.name, "Source Path": str(f), "Sheet Name": info["sheet"],
                "Issue Type": "EHT_COLUMN_CHECK", "Message": problem,
                "Row Snapshot": " | ".join(info["labels"]),
            })
    target_files = [f for f in target_files if f not in eht_files]

    print(f"검색된 xlsx 파일 수: {len(files)}")
    print(f"실제 처리 대상 파일 수: {len(target_files)}")

    if not target_files:
        print("처리할 .xlsx 파일이 없습니다.")
        input("엔터를 누르면 종료합니다...")
        return

    for items, excluded, summary, parse_issues in process_files(sorted(target_files)):
        all_items.extend(items)
        all_excluded.extend(excluded)
        summaries.append(summary)
        errors.extend(parse_issues)

    eht_sheets = []
    eht_summary = []
    if eht_files:
        print("EHT LIST 비교 중...")
        eht_report, eht_summary = compare_eht(eht_rows, all_items, eht_infos)
        eht_sheets = [
            ("EHT비교요약", eht_summary, EHT_SUMMARY_COLUMNS),
            ("EHT비교", eht_report, EHT_REPORT_COLUMNS),
        ]

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = base_dir / f"INSULATION_통합정리_{timestamp}.xlsx"

    print("결과 파일 저장 중...")
    save_output_excel(
        output_path=output_path,
        all_items=all_items,
        all_excluded=all_excluded,
        summaries=summaries,
        errors=errors,
        eht_sheets=eht_sheets,
    )

    print("\n=== 작업 완료 ===")
    print(f"통합 파일: {output_path}")
    print(f"총 아이템 수: {len(all_items)}")
    print(f"제외라인 수: {len(all_excluded)}")
    print(f"에러/검토이슈 수: {len(errors)}")
    if eht_summary:
        print("EHT 비교: " + ", ".join(f"{rec['판정']} {rec['건수']}" for rec in eht_summary if rec.get("건수")))
    print(f"소요 시간: {time.perf_counter() - started:.1f}초")

    input("엔터를 누르면 종료합니다...")


if __name__ == "__main__":
    freeze_support()  # exe로 묶어 실행할 때 병렬 처리용
    try:
        main()
    except Exception as e:
        print("### 프로그램 치명적 오류 ###")
        print(e)
        print(traceback.format_exc())
        input("엔터를 누르면 종료합니다...")
